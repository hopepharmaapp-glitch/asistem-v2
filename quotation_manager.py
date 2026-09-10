"""
📋 Quotation / Proposal Manager for HopePharma
=================================================
• Dedicated creation UI with all metadata fields
• List view with search + filters + status badges
• Manual item add (Product, Description, Qty, Price, Total, Taxable)
• Full input validation
• Persistent JSON storage via EnhancedCloudDataManager
• Convert Quotation → Sales Invoice in 1 click
• PDF/Excel export via reportlab + openpyxl
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, date, timedelta
from decimal import Decimal
import os
import sys

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
if _THIS_DIR not in sys.path:
    sys.path.insert(0, _THIS_DIR)

try:
    from hope_pharma_complete import QuotationObject
except Exception:
    QuotationObject = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except Exception:
    HAS_OPENPYXL = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    HAS_REPORTLAB = True
except Exception:
    HAS_REPORTLAB = False


# ===========================================================================
# 1. Popup: Create / Edit ONE Quotation
# ===========================================================================
class QuotationFormDialog:
    def __init__(self, parent, data_manager, existing_quote=None, on_saved_cb=None):
        self.dm = data_manager
        self.quote = existing_quote  # None = new
        self.on_saved = on_saved_cb
        self.dlg = tk.Toplevel(parent)
        self.dlg.title(("✏️ Edit Quotation" if existing_quote else "➕ New Quotation / Proposal"))
        self.dlg.geometry("960x740")
        self.dlg.minsize(860, 640)
        self.dlg.transient(parent)
        self.items = list((existing_quote or {}).get("items", []))  # mutable copy
        self._build()

    # ---------- UI ----------
    def _build(self):
        root = ttk.Frame(self.dlg, padding=14)
        root.pack(fill="both", expand=True)

        # -------- Metadata --------
        meta = ttk.LabelFrame(root, text="📋 Quotation Details", padding=12)
        meta.pack(fill="x")
        meta.columnconfigure(1, weight=1)
        meta.columnconfigure(3, weight=1)

        qdata = self.quote or {}
        today = date.today().isoformat()
        valid = (date.today() + timedelta(days=30)).isoformat()

        fields = [
            ("Quotation ID:", 0, "quotation_id", qdata.get("quotation_id", "") or self._next_id(), True),
            ("Date:", 0, "date", qdata.get("date", today), False),
            ("Valid Until:", 0, "valid_until", qdata.get("valid_until", valid), False),
            ("Client Name *:", 1, "client_name", qdata.get("client_name", ""), False),
            ("Client TRN:", 1, "client_trn", qdata.get("client_trn", ""), False),
            ("Client Emirate:", 2, "client_emirate", qdata.get("client_emirate", "Dubai"), False),
            ("Currency:", 2, "currency", qdata.get("currency", "AED"), False),
            ("VAT Rate %:", 3, "tax_rate", qdata.get("tax_rate", 5), False),
            ("Status:", 3, "status", qdata.get("status", "Draft"), False),
        ]
        self.vars = {}
        for label, row, key, val, ro in fields:
            col = 1 if row in (0, 2) else 3
            if row >= 2:
                r = row - 2
            else:
                r = row
            col_lbl = col - 1
            if key == "status":
                opts = ["Draft", "Sent", "Accepted", "Rejected", "Expired", "Converted to Invoice"]
            elif key == "currency":
                opts = ["AED", "USD", "EUR", "GBP", "SAR", "QAR", "OMR", "KWD", "BHD", "EGP"]
            elif key == "client_emirate":
                opts = ["Dubai", "Abu Dhabi", "Sharjah", "Ajman",
                        "Umm Al Quwain", "Ras Al Khaimah", "Fujairah", "Other"]
            else:
                opts = None
            ttk.Label(meta, text=label,
                      font=("Helvetica", 9, "bold")).grid(row=r, column=col_lbl, sticky="e", padx=(0, 6), pady=4)
            v = tk.StringVar(value=str(val))
            self.vars[key] = v
            if ro:
                en = ttk.Entry(meta, textvariable=v, state="readonly", width=24)
            elif opts:
                en = ttk.Combobox(meta, textvariable=v, values=opts,
                                  state="readonly" if key in ("status", "currency", "client_emirate") else "normal",
                                  width=24)
            else:
                en = ttk.Entry(meta, textvariable=v, width=24)
            en.grid(row=r, column=col, sticky="we", pady=4, padx=(0, 12))

        # Notes
        ttk.Label(meta, text="Notes / Terms:",
                  font=("Helvetica", 9, "bold")).grid(row=3, column=0, sticky="ne", padx=(0, 6), pady=4)
        self.notes_txt = tk.Text(meta, height=3, wrap="word", relief="solid", borderwidth=1,
                                  font=("Helvetica", 9))
        self.notes_txt.grid(row=3, column=1, columnspan=3, sticky="we", pady=4)
        self.notes_txt.insert("1.0", qdata.get("notes", ""))

        # -------- Items --------
        items_lf = ttk.LabelFrame(root, text="🧾 Line Items", padding=10)
        items_lf.pack(fill="both", expand=True, pady=(12, 0))

        # Form to add item
        addfrm = ttk.Frame(items_lf)
        addfrm.pack(fill="x", pady=(0, 8))

        def _lbl(x):
            ttk.Label(addfrm, text=x, font=("Helvetica", 9, "bold")).pack(side="left", padx=(6, 3))

        _lbl("Product / Service:")
        self.f_product = ttk.Entry(addfrm, width=26, font=("Helvetica", 9))
        self.f_product.pack(side="left", padx=(0, 6))
        _lbl("Description:")
        self.f_desc = ttk.Entry(addfrm, width=36, font=("Helvetica", 9))
        self.f_desc.pack(side="left", padx=(0, 6))
        _lbl("Qty:")
        self.f_qty = ttk.Entry(addfrm, width=8, font=("Helvetica", 9))
        self.f_qty.insert(0, "1")
        self.f_qty.pack(side="left", padx=(0, 6))
        _lbl("Unit Price:")
        self.f_price = ttk.Entry(addfrm, width=12, font=("Helvetica", 9))
        self.f_price.insert(0, "0.00")
        self.f_price.pack(side="left", padx=(0, 6))
        self.f_taxable = tk.BooleanVar(value=True)
        ttk.Checkbutton(addfrm, text="Taxable", variable=self.f_taxable).pack(side="left", padx=(0, 10))
        ttk.Button(addfrm, text="➕ Add Line Item", command=self._add_item,
                   style="Accent.TButton").pack(side="left")
        ttk.Button(addfrm, text="➖ Remove Selected", command=self._remove_item).pack(side="left", padx=6)

        cols = ("idx", "product", "description", "qty", "unit_price", "total", "taxable")
        self.item_tree = ttk.Treeview(items_lf, columns=cols, show="headings", height=10)
        for c, lbl, w, a in zip(cols,
                                 ("#", "Product / Service", "Description", "Qty", "Unit Price", "Line Total", "Taxable"),
                                 (40, 180, 260, 60, 100, 120, 80),
                                 ("center", "w", "w", "e", "e", "e", "center")):
            self.item_tree.heading(c, text=lbl)
            self.item_tree.column(c, width=w, anchor=a, stretch=(c == "description"))
        self.item_tree.pack(fill="both", expand=True)
        for i, it in enumerate(self.items, start=1):
            self._push_item_row(i, it)

        # -------- Totals --------
        totals = ttk.Frame(root)
        totals.pack(fill="x", pady=(10, 0))
        self.totals_label = ttk.Label(totals, text="(0 items)", font=("Helvetica", 10, "bold"),
                                      foreground="#1A365D", anchor="e", justify="right")
        self.totals_label.pack(fill="x")
        self._recalc_totals_view()

        # -------- Buttons --------
        btns = ttk.Frame(root)
        btns.pack(fill="x", pady=(14, 0))
        ttk.Button(btns, text="❌ Cancel", command=self.dlg.destroy).pack(side="right", padx=(6, 0))
        ttk.Button(btns, text="💾 Save Quotation", command=self._save,
                   style="Accent.TButton").pack(side="right")

    # ---------- Helpers ----------
    def _next_id(self):
        try:
            settings = self.dm.load_json("app_settings.json") or {}
            prefix = settings.get("quote_prefix", "QUO-")
            last = int(settings.get("last_quote_number", 0) or 0) + 1
            settings["last_quote_number"] = last
            self.dm.save_json("app_settings.json", settings)
            return f"{prefix}{date.today().strftime('%Y%m')}-{last:04d}"
        except Exception:
            return f"QUO-{datetime.now().strftime('%Y%m%d%H%M%S')}"

    def _add_item(self):
        prod = self.f_product.get().strip()
        desc = self.f_desc.get().strip()
        if not (prod or desc):
            messagebox.showwarning("Invalid Line",
                                   "Product or Description is required for each line item.")
            return
        try:
            qty = int(self.f_qty.get())
            if qty <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Qty", "Quantity must be a positive integer.")
            return
        try:
            price = float(self.f_price.get())
            if price < 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Invalid Price", "Unit Price must be a non-negative number.")
            return
        line = {
            "product": prod,
            "description": desc,
            "quantity": qty,
            "unit_price": round(price, 2),
            "total": round(qty * price, 2),
            "taxable": bool(self.f_taxable.get()),
        }
        self.items.append(line)
        self._push_item_row(len(self.items), line)
        self._recalc_totals_view()
        for f in (self.f_product, self.f_desc):
            f.delete(0, "end")
        self.f_qty.delete(0, "end"); self.f_qty.insert(0, "1")
        self.f_price.delete(0, "end"); self.f_price.insert(0, "0.00")

    def _push_item_row(self, idx, it):
        self.item_tree.insert("", "end",
                              values=(idx, it.get("product", ""), it.get("description", ""),
                                      f"{it.get('quantity', 1):,}",
                                      f"{float(it.get('unit_price', 0)):,.2f}",
                                      f"{float(it.get('total', 0)):,.2f}",
                                      ("Yes" if it.get("taxable", True) else "No")))

    def _remove_item(self):
        sel = self.item_tree.selection()
        if not sel:
            return
        for s in reversed(sel):
            idx = self.item_tree.index(s)
            if 0 <= idx < len(self.items):
                self.items.pop(idx)
        self.item_tree.delete(*self.item_tree.get_children())
        for i, it in enumerate(self.items, start=1):
            self._push_item_row(i, it)
        self._recalc_totals_view()

    def _recalc_totals_view(self):
        subtotal = sum(float(i.get("total", 0)) for i in self.items)
        taxable = sum(float(i.get("total", 0)) for i in self.items if i.get("taxable", True))
        try:
            rate = float(self.vars["tax_rate"].get())
        except Exception:
            rate = 0.0
        tax = round(taxable * (rate / 100.0), 2)
        total = round(subtotal + tax, 2)
        curr = self.vars["currency"].get()
        self.totals_label.config(
            text=(f"Subtotal: {curr} {subtotal:,.2f}    |    "
                  f"Taxable Amount: {curr} {taxable:,.2f}  ×  {rate}%  =  {curr} {tax:,.2f}    |    "
                  f"👉 GRAND TOTAL: {curr} {total:,.2f}    ({len(self.items)} line items)")
        )

    # ---------- Save ----------
    def _save(self):
        if not self.vars["client_name"].get().strip():
            messagebox.showwarning("Missing Client", "Client Name is required.")
            return
        if len(self.items) == 0:
            if not messagebox.askyesno("No Line Items",
                                       "This quotation has no items. Save anyway?"):
                return
        subtotal = sum(float(i.get("total", 0)) for i in self.items)
        taxable = sum(float(i.get("total", 0)) for i in self.items if i.get("taxable", True))
        try:
            rate = float(self.vars["tax_rate"].get())
        except Exception:
            rate = 0.0
        tax = round(taxable * (rate / 100.0), 2)
        payload = {
            "quotation_id": self.vars["quotation_id"].get().strip(),
            "date": self.vars["date"].get().strip(),
            "valid_until": self.vars["valid_until"].get().strip(),
            "client_name": self.vars["client_name"].get().strip(),
            "client_trn": self.vars["client_trn"].get().strip(),
            "client_emirate": self.vars["client_emirate"].get().strip(),
            "currency": self.vars["currency"].get().strip(),
            "tax_rate": rate,
            "items": self.items,
            "subtotal": round(subtotal, 2),
            "taxable_amount": round(taxable, 2),
            "non_taxable_amount": round(subtotal - taxable, 2),
            "tax_amount": tax,
            "grand_total": round(subtotal + tax, 2),
            "notes": self.notes_txt.get("1.0", "end").strip(),
            "status": self.vars["status"].get().strip() or "Draft",
            "created_at": (self.quote or {}).get("created_at", datetime.now().isoformat()),
            "updated_at": datetime.now().isoformat(),
        }
        if QuotationObject:
            try:
                payload = QuotationObject(payload).to_dict()
            except Exception:
                pass
        quotes = self.dm.load_json("quotes_data.json") or []
        replaced = False
        for i, q in enumerate(quotes):
            if q.get("quotation_id") == payload["quotation_id"]:
                quotes[i] = payload; replaced = True; break
        if not replaced:
            quotes.append(payload)
        if not self.dm.save_json("quotes_data.json", quotes):
            messagebox.showerror("Error", "Could not save quotation data.")
            return
        messagebox.showinfo("✅ Saved",
                            f"Quotation {payload['quotation_id']} saved.\n\n"
                            f"Client: {payload['client_name']}\n"
                            f"Grand Total: {payload['currency']} {payload['grand_total']:,.2f}\n"
                            f"Valid until: {payload['valid_until']}")
        if callable(self.on_saved):
            try:
                self.on_saved()
            except Exception:
                pass
        self.dlg.destroy()


# ===========================================================================
# 2. Manager: List, Search, Export, Convert → Invoice
# ===========================================================================
class QuotationManagerDialog:
    def __init__(self, parent, data_manager, invoice_app=None):
        self.dm = data_manager
        self.app = invoice_app  # optional, for convert-to-invoice callback
        self.win = tk.Toplevel(parent)
        self.win.title("📋 Quotation / Proposal Manager — HopePharma")
        self.win.geometry("1260x760")
        self.win.minsize(1000, 600)
        self.win.transient(parent)
        self._style()
        self._build()
        self._refresh()

    def _style(self):
        s = ttk.Style(self.win)
        try: s.theme_use("clam")
        except Exception: pass
        s.configure("Accent.TButton", font=("Helvetica", 10, "bold"))
        s.map("Accent.TButton",
              background=[("active", "#3182CE"), ("!disabled", "#1A365D")],
              foreground=[("!disabled", "white")])

    def _build(self):
        header = ttk.Frame(self.win, padding=(16, 12))
        header.pack(fill="x")
        ttk.Label(header, text="📋 Quotation / Proposal Manager",
                  font=("Helvetica", 20, "bold"),
                  foreground="#1A365D").pack(anchor="w")
        ttk.Label(header,
                  text="Create professional quotes, send to clients, then convert to sales invoices in one click.",
                  font=("Helvetica", 9), foreground="#4A5568").pack(anchor="w")

        # Toolbar
        bar = ttk.Frame(self.win, padding=(16, 0, 16, 8))
        bar.pack(fill="x")
        ttk.Button(bar, text="➕ New Quotation",
                   command=self._new, style="Accent.TButton").pack(side="left")
        ttk.Button(bar, text="✏️ Edit Selected", command=self._edit).pack(side="left", padx=6)
        ttk.Button(bar, text="❌ Delete Selected", command=self._delete).pack(side="left", padx=6)
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=10)
        ttk.Button(bar, text="🧾 Convert to Invoice", command=self._convert).pack(side="left", padx=6)
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=10)
        ttk.Button(bar, text="📄 Export PDF", command=lambda: self._export("pdf")).pack(side="left", padx=6)
        ttk.Button(bar, text="📊 Export Excel", command=lambda: self._export("xlsx")).pack(side="left", padx=6)

        # Search
        search = ttk.Frame(self.win, padding=(16, 0, 16, 8))
        search.pack(fill="x")
        ttk.Label(search, text="🔍 Search:",
                  font=("Helvetica", 9, "bold")).pack(side="left")
        self.s_q = tk.StringVar()
        e = ttk.Entry(search, textvariable=self.s_q, width=48)
        e.pack(side="left", padx=6)
        e.bind("<KeyRelease>", lambda _e: self._refresh())
        ttk.Label(search, text="Status:",
                  font=("Helvetica", 9, "bold")).pack(side="left", padx=(16, 0))
        self.s_status = tk.StringVar(value="All")
        ttk.Combobox(search, textvariable=self.s_status, width=18, state="readonly",
                     values=["All", "Draft", "Sent", "Accepted", "Rejected", "Expired",
                             "Converted to Invoice"]).pack(side="left", padx=6)
        self.s_status.trace_add("write", lambda *_: self._refresh())
        ttk.Button(search, text="🔄 Refresh", command=self._refresh).pack(side="right")

        cols = ("id", "date", "valid_until", "client", "status",
                "currency", "subtotal", "tax", "grand_total", "items")
        self.tree = ttk.Treeview(self.win, columns=cols, show="headings", height=20)
        widths = (140, 100, 100, 260, 150, 80, 120, 100, 130, 70)
        aligns = ("w", "w", "w", "w", "w", "center", "e", "e", "e", "center")
        for c, lbl, w, a in zip(cols,
                                 ("Quotation #", "Date", "Valid Until", "Client", "Status",
                                  "Curr", "Subtotal", "VAT", "Grand Total", "# Items"),
                                 widths, aligns):
            self.tree.heading(c, text=lbl)
            self.tree.column(c, width=w, anchor=a, stretch=(c == "client"))
        # Colour tags
        for tag, fg in [("Accepted", "#38A169"), ("Rejected", "#E53E3E"),
                        ("Expired", "#B7791F"),
                        ("Converted to Invoice", "#2B6CB0"),
                        ("Sent", "#805AD5"), ("Draft", "#4A5568")]:
            self.tree.tag_configure(tag, foreground=fg, font=("Helvetica", 9, "bold"))
        self.tree.pack(fill="both", expand=True, padx=16, pady=(0, 4))
        self.tree.bind("<Double-1>", lambda _e: self._edit())

        # Footer summary
        self.footer = ttk.Label(self.win, text="", padding=(16, 8),
                                font=("Helvetica", 9, "bold"), foreground="#1A365D")
        self.footer.pack(fill="x")

    # ---------- Data ----------
    def _all_quotes(self):
        return self.dm.load_json("quotes_data.json") or []

    def _refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = (self.s_q.get() or "").strip().lower()
        st = self.s_status.get()
        rows = []
        for x in self._all_quotes():
            if st != "All" and x.get("status", "") != st:
                continue
            if q:
                hay = " ".join([str(x.get(k, "")) for k in
                                ("quotation_id", "client_name", "client_trn",
                                 "status", "notes")]).lower()
                if q not in hay:
                    continue
            rows.append(x)
        rows.sort(key=lambda r: r.get("date", ""), reverse=True)
        total_val = Decimal("0")
        curr_set = set()
        for r in rows:
            curr = r.get("currency", "AED")
            curr_set.add(curr)
            if len(curr_set) == 1:
                total_val += Decimal(str(r.get("grand_total", 0)))
            vals = (r.get("quotation_id", ""),
                    r.get("date", ""),
                    r.get("valid_until", ""),
                    r.get("client_name", ""),
                    r.get("status", ""),
                    curr,
                    f"{float(r.get('subtotal', 0)):,.2f}",
                    f"{float(r.get('tax_amount', 0)):,.2f}",
                    f"{float(r.get('grand_total', 0)):,.2f}",
                    str(len(r.get("items", []))))
            tag = (r.get("status", "Draft"),)
            self.tree.insert("", "end", values=vals, tags=tag)
        # Footer
        if len(curr_set) == 1 and rows:
            footer_text = f"Showing {len(rows)} quotation(s)  |  Total value {curr_set.pop()} {float(total_val):,.2f}"
        else:
            footer_text = f"Showing {len(rows)} quotation(s)  |  Total value: mixed currencies (open individually)"
        counts = {}
        for r in self._all_quotes():
            counts[r.get("status", "Draft")] = counts.get(r.get("status", "Draft"), 0) + 1
        if counts:
            footer_text += "    |    Breakdown:  " + "  ·  ".join(f"{k}: {v}" for k, v in counts.items())
        self.footer.config(text=footer_text)

    # ---------- Actions ----------
    def _selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Select a Quotation",
                                "Please select one quotation from the list first.")
            return None
        id_ = self.tree.item(sel[0], "values")[0]
        for q in self._all_quotes():
            if q.get("quotation_id") == id_:
                return q
        return None

    def _new(self):
        QuotationFormDialog(self.win, self.dm, on_saved_cb=self._refresh)

    def _edit(self):
        q = self._selected()
        if not q: return
        QuotationFormDialog(self.win, self.dm, existing_quote=q, on_saved_cb=self._refresh)

    def _delete(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Select", "Please select quotation(s) to delete.")
            return
        if not messagebox.askyesno("Confirm",
                                   f"Delete {len(sel)} quotation(s)? This cannot be undone."):
            return
        to_del = {self.tree.item(s, "values")[0] for s in sel}
        new = [q for q in self._all_quotes() if q.get("quotation_id") not in to_del]
        self.dm.save_json("quotes_data.json", new)
        self._refresh()

    def _convert(self):
        q = self._selected()
        if not q: return
        if q.get("status") == "Converted to Invoice":
            if not messagebox.askyesno("Already Converted",
                                       "This quote was already converted. Create another invoice?"):
                return
        items = []
        for it in q.get("items", []):
            items.append({
                "id": "", "product_id": "",
                "description": (it.get("product", "") + " — " + it.get("description", "")).strip(" —") or it.get("product", "") or "Item",
                "quantity": it.get("quantity", 1),
                "unit_price": float(it.get("unit_price", 0)),
                "total": float(it.get("total", 0)),
                "taxable": it.get("taxable", True),
                "cost": 0,
            })
        inv = {
            "invoice_id": "",  # generated by manager
            "date": date.today().isoformat(),
            "due_date": (date.today() + timedelta(days=30)).isoformat(),
            "client_name": q.get("client_name", ""),
            "client_trn": q.get("client_trn", ""),
            "client_emirate": q.get("client_emirate", ""),
            "client_location": "",
            "currency": q.get("currency", "AED"),
            "tax_rate": float(q.get("tax_rate", 5)),
            "items": items,
            "subtotal": float(q.get("subtotal", 0)),
            "taxable_amount": float(q.get("taxable_amount", 0)),
            "non_taxable_amount": float(q.get("non_taxable_amount", 0)),
            "tax_amount": float(q.get("tax_amount", 0)),
            "grand_total": float(q.get("grand_total", 0)),
            "total_cost": 0,
            "profit": float(q.get("grand_total", 0)),
            "total_paid": 0,
            "balance_due": float(q.get("grand_total", 0)),
            "status": "Not Paid",
            "payment_history": [],
            "costs": [],
            "notes": f"Converted from Quotation {q.get('quotation_id', '')}\n" + (q.get("notes") or ""),
            "source_quotation": q.get("quotation_id", ""),
        }
        try:
            inv["invoice_id"] = self.dm.generate_invoice_id()
        except Exception:
            inv["invoice_id"] = f"INV-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        if messagebox.askyesno("Create Invoice?",
                               f"Create invoice {inv['invoice_id']} for {inv['client_name']} — "
                               f"{inv['currency']} {inv['grand_total']:,.2f}?"):
            ok = self.dm.add_invoice_from_dict(inv)
            if not ok:
                messagebox.showerror("Failed", "Could not create invoice.")
                return
            # Mark quotation as converted
            quotes = self._all_quotes()
            for i, x in enumerate(quotes):
                if x.get("quotation_id") == q.get("quotation_id"):
                    quotes[i]["status"] = "Converted to Invoice"
                    quotes[i]["converted_to_invoice"] = inv["invoice_id"]
                    quotes[i]["converted_at"] = datetime.now().isoformat()
            self.dm.save_json("quotes_data.json", quotes)
            messagebox.showinfo("✅ Done",
                                f"Invoice {inv['invoice_id']} created.\n\n"
                                f"Quotation status updated to 'Converted to Invoice'.")
            self._refresh()

    # ---------- Exports ----------
    def _export(self, fmt):
        q = self._selected()
        if not q: return
        today = date.today().isoformat()
        base = f"HopePharma_Quote_{q.get('quotation_id','').replace('/','-')}_{today}"
        if fmt == "xlsx":
            if not HAS_OPENPYXL:
                messagebox.showerror("Missing", "pip install openpyxl")
                return
            out = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel", "*.xlsx")],
                initialfile=base + ".xlsx")
            if not out: return
            self._export_xlsx(q, out)
        else:
            if not HAS_REPORTLAB:
                messagebox.showerror("Missing", "pip install reportlab")
                return
            out = filedialog.asksaveasfilename(
                defaultextension=".pdf",
                filetypes=[("PDF", "*.pdf")],
                initialfile=base + ".pdf")
            if not out: return
            self._export_pdf(q, out)
        if messagebox.askyesno("Exported", f"Saved to:\n{out}\n\nOpen now?"):
            try:
                import platform, subprocess
                if platform.system() == "Darwin": subprocess.Popen(["open", out])
                elif platform.system() == "Windows": os.startfile(out)
                else: subprocess.Popen(["xdg-open", out])
            except Exception:
                pass

    def _export_xlsx(self, q, out):
        wb = Workbook()
        ws = wb.active; ws.title = "Quotation"
        styles = {
            "h": Font(bold=True, color="FFFFFF", size=12),
            "hf": PatternFill("solid", fgColor="1A365D"),
            "t": Font(bold=True, size=16, color="1A365D"),
            "tf": Font(bold=True, size=11, color="1A365D"),
            "total": Font(bold=True, size=11),
            "totalf": PatternFill("solid", fgColor="EBF8FF"),
            "c": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "r": Alignment(horizontal="right"),
            "l": Alignment(horizontal="left", wrap_text=True),
            "b": Border(
                left=Side(style="thin", color="CBD5E0"),
                right=Side(style="thin", color="CBD5E0"),
                top=Side(style="thin", color="CBD5E0"),
                bottom=Side(style="thin", color="CBD5E0")),
        }
        ws.merge_cells("A1:F1")
        ws["A1"] = f"🏥 HopePharma — QUOTATION / PROFORMA INVOICE"
        ws["A1"].font = styles["t"]
        ws.merge_cells("A2:F2")
        ws["A2"] = f"Quotation # {q.get('quotation_id','')}   |   Date: {q.get('date','')}   |   Valid Until: {q.get('valid_until','')}   |   Status: {q.get('status','')}"
        ws["A2"].font = Font(italic=True, color="7F8C8D")
        # Client
        ws.merge_cells("A4:B4"); ws["A4"] = "Client Details:"; ws["A4"].font = styles["tf"]
        info = [
            ("Client:", q.get("client_name", "")),
            ("TRN:", q.get("client_trn", "")),
            ("Emirate:", q.get("client_emirate", "")),
            ("Currency:", q.get("currency", "AED")),
        ]
        for i, (k, v) in enumerate(info, start=5):
            ws.cell(row=i, column=1, value=k).font = styles["total"]
            ws.cell(row=i, column=2, value=v)
        # Items
        r = 11
        headers = ["#", "Product / Service", "Description", "Qty", "Unit Price", "Line Total"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = styles["h"]; cell.fill = styles["hf"]
            cell.alignment = styles["c"]; cell.border = styles["b"]
        r += 1
        for i, it in enumerate(q.get("items", []), start=1):
            for c, val in enumerate([
                i,
                it.get("product", ""),
                it.get("description", ""),
                it.get("quantity", 1),
                float(it.get("unit_price", 0)),
                float(it.get("total", 0)),
            ], 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.border = styles["b"]
                if c in (4, 5, 6):
                    cell.alignment = styles["r"]
                    if isinstance(val, float): cell.number_format = '#,##0.00'
                elif c == 1: cell.alignment = styles["c"]
                else: cell.alignment = styles["l"]
            r += 1
        # Totals
        curr = q.get("currency", "AED")
        totals = [
            ("Subtotal", q.get("subtotal", 0)),
            (f"VAT ({q.get('tax_rate', 5)}%)", q.get("tax_amount", 0)),
            ("GRAND TOTAL", q.get("grand_total", 0)),
        ]
        for label, val in totals:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            c1 = ws.cell(row=r, column=1, value=label)
            c1.alignment = Alignment(horizontal="right"); c1.font = styles["total"]
            if "GRAND" in label: c1.fill = styles["totalf"]
            c2 = ws.cell(row=r, column=6, value=float(val or 0)); c2.number_format = '#,##0.00'
            c2.font = styles["total"]; c2.alignment = styles["r"]
            if "GRAND" in label: c2.fill = styles["totalf"]
            r += 1
        if q.get("notes"):
            r += 2
            ws.cell(row=r, column=1, value="Notes / Terms:").font = styles["tf"]
            r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r + 4, end_column=6)
            ws.cell(row=r, column=1, value=q["notes"]).alignment = styles["l"]
        for col, w in zip("ABCDEF", [6, 28, 42, 8, 14, 16]):
            ws.column_dimensions[col].width = w
        wb.save(out)

    def _export_pdf(self, q, out):
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
        from reportlab.lib.units import cm
        doc = SimpleDocTemplate(out, pagesize=A4,
                                rightMargin=2*cm, leftMargin=2*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        t = ParagraphStyle("T", parent=styles["Title"], textColor=colors.HexColor("#1A365D"),
                           alignment=TA_CENTER, fontSize=20, spaceAfter=6)
        sub = ParagraphStyle("S", parent=styles["Normal"], textColor=colors.HexColor("#7F8C8D"),
                             alignment=TA_CENTER, fontSize=9, spaceAfter=12)
        h3 = ParagraphStyle("H3", parent=styles["Heading3"], textColor=colors.HexColor("#1A365D"), fontSize=12, spaceAfter=4)
        story = []
        story.append(Paragraph("🏥 HopePharma Medical Trading L.L.C.", t))
        story.append(Paragraph(f"<b>QUOTATION / PROFORMA INVOICE</b> — {q.get('quotation_id','')}", t))
        story.append(Paragraph(f"Date: <b>{q.get('date','')}</b> | Valid Until: <b>{q.get('valid_until','')}</b> | Status: <b>{q.get('status','')}</b>", sub))
        # Client block
        tbl = [
            ["Client:", q.get("client_name", ""), "Currency:", q.get("currency", "AED")],
            ["TRN:", q.get("client_trn", ""), "VAT Rate:", f"{q.get('tax_rate', 5)}%"],
            ["Emirate:", q.get("client_emirate", ""), "", ""],
        ]
        pt = Table(tbl, colWidths=[2.5*cm, 6.5*cm, 2.5*cm, 5.5*cm], hAlign="LEFT")
        pt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#A0AEC0")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAFC")),
        ]))
        story.append(pt); story.append(Spacer(1, 0.6*cm))
        story.append(Paragraph("🧾 Line Items", h3))
        header = ["#", "Product / Service", "Description", "Qty", "Unit Price", "Total"]
        data = [header]
        for i, it in enumerate(q.get("items", []), start=1):
            data.append([
                str(i),
                Paragraph(str(it.get("product", "")), styles["Normal"]),
                Paragraph(str(it.get("description", "")), styles["Normal"]),
                f"{it.get('quantity',1):,}",
                f"{float(it.get('unit_price', 0)):,.2f}",
                f"{float(it.get('total', 0)):,.2f}",
            ])
        items_tbl = Table(data, repeatRows=1,
                          colWidths=[0.8*cm, 3.6*cm, 5.4*cm, 1.2*cm, 2.5*cm, 2.5*cm])
        items_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#A0AEC0")),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("ALIGN", (3, 1), (5, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F7FAFC")]),
        ]))
        story.append(items_tbl); story.append(Spacer(1, 0.5*cm))
        # Totals
        curr = q.get("currency", "AED")
        totals = [
            ["Subtotal", f"{curr} {float(q.get('subtotal',0)):,.2f}"],
            [f"VAT ({q.get('tax_rate',5)}%)", f"{curr} {float(q.get('tax_amount',0)):,.2f}"],
            ["GRAND TOTAL", f"{curr} {float(q.get('grand_total',0)):,.2f}"],
        ]
        tt = Table(totals, colWidths=[13*cm, 3*cm], hAlign="RIGHT")
        tt.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -2), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#A0AEC0")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#C6F6D5")),
            ("BACKGROUND", (0, 0), (-1, -2), colors.HexColor("#EBF8FF")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(tt)
        if q.get("notes"):
            story.append(Spacer(1, 0.6*cm))
            story.append(Paragraph("📝 Notes / Terms & Conditions", h3))
            story.append(Paragraph(str(q["notes"]).replace("\n", "<br/>"), styles["Normal"]))
        story.append(Spacer(1, 1*cm))
        story.append(Paragraph(
            "<i>This quotation is valid until the date stated above. E.&amp;O.E.</i>",
            ParagraphStyle("foot", parent=styles["Normal"], fontSize=8,
                           textColor=colors.HexColor("#7F8C8D"), alignment=TA_CENTER)))
        doc.build(story)
