import os
import json
import csv
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox
import webbrowser
import getpass

from inventory_system import InventoryManager, build_pdf_logo_flowables
from warehouse_extension import WarehouseManager

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False


def getpass_user_fallback():
    try:
        return getpass.getuser()
    except Exception:
        return "System User"


class InventoryReportScreen:
    def __init__(self, parent, invoice_manager, back_callback=None):
        self.parent = parent
        self.invoice_manager = invoice_manager
        self.back_callback = back_callback
        self.data_folder = getattr(invoice_manager, "invoice_folder", os.getcwd())
        self.inventory_manager = InventoryManager(self.data_folder)
        self.search_rows = []
        self.goods_rows = []
        self.inbound_rows = []
        self.outbound_rows = []
        self.outbound_filtered = []
        self.outbound_page = 1
        self.outbound_page_size = 25
        self._build_ui()
        self.refresh_all()

    def _build_ui(self):
        top = ttk.Frame(self.parent)
        top.pack(fill="x", pady=(0, 8))

        if self.back_callback:
            ttk.Button(top, text="← Back to Main Menu", command=self.back_callback).pack(side="left")
        ttk.Label(top, text="Inventory Report", font=("Helvetica", 20, "bold")).pack(side="left", padx=12)

        action_bar = ttk.Frame(self.parent)
        action_bar.pack(fill="x", pady=(0, 8))
        ttk.Button(action_bar, text="Inventory Search", command=lambda: self.notebook.select(self.search_tab)).pack(side="left")
        ttk.Button(action_bar, text="Goods Received", command=lambda: self.notebook.select(self.goods_tab)).pack(side="left", padx=6)
        ttk.Button(action_bar, text="Inbound Report", command=lambda: self.notebook.select(self.inbound_tab)).pack(side="left")
        ttk.Button(action_bar, text="Outbound Report", command=lambda: self.notebook.select(self.outbound_tab)).pack(side="left", padx=6)
        ttk.Button(action_bar, text="Refresh", command=self.refresh_all).pack(side="right")

        self.notebook = ttk.Notebook(self.parent)
        self.notebook.pack(fill="both", expand=True)

        self.search_tab = ttk.Frame(self.notebook, padding=10)
        self.goods_tab = ttk.Frame(self.notebook, padding=10)
        self.inbound_tab = ttk.Frame(self.notebook, padding=10)
        self.outbound_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.search_tab, text="Inventory Search")
        self.notebook.add(self.goods_tab, text="Goods Received")
        self.notebook.add(self.inbound_tab, text="Inbound Report")
        self.notebook.add(self.outbound_tab, text="Outbound Report")

        self._build_search_tab()
        self._build_goods_tab()
        self._build_inbound_tab()
        self._build_outbound_tab()

        self.status_var = tk.StringVar(value="Ready")
        ttk.Label(self.parent, textvariable=self.status_var, relief=tk.SUNKEN, anchor="w").pack(fill="x", pady=(8, 0))

    def _build_search_tab(self):
        filters = ttk.LabelFrame(self.search_tab, text="Search Filters", padding=10)
        filters.pack(fill="x", pady=(0, 8))

        self.code_var = tk.StringVar()
        self.product_var = tk.StringVar()
        self.batch_var = tk.StringVar()
        self.manufacture_var = tk.StringVar()
        self.expiry_var = tk.StringVar()
        self.client_var = tk.StringVar()

        ttk.Label(filters, text="Code").grid(row=0, column=0, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.code_var, width=18).grid(row=0, column=1, sticky="w", padx=(0, 12), pady=4)
        ttk.Label(filters, text="Product Name").grid(row=0, column=2, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.product_var, width=24).grid(row=0, column=3, sticky="w", padx=(0, 12), pady=4)
        ttk.Label(filters, text="Batch Number").grid(row=1, column=0, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.batch_var, width=18).grid(row=1, column=1, sticky="w", padx=(0, 12), pady=4)
        ttk.Label(filters, text="Manufacture Date").grid(row=1, column=2, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.manufacture_var, width=24).grid(row=1, column=3, sticky="w", padx=(0, 12), pady=4)
        ttk.Label(filters, text="Expiry Date").grid(row=2, column=0, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.expiry_var, width=18).grid(row=2, column=1, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(filters, text="Client Name").grid(row=2, column=2, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(filters, textvariable=self.client_var, width=24).grid(row=2, column=3, sticky="w", padx=(0, 12), pady=4)

        btns = ttk.Frame(filters)
        btns.grid(row=3, column=3, sticky="e", pady=4)
        ttk.Button(btns, text="Search", command=self.run_search).pack(side="left")
        ttk.Button(btns, text="Clear", command=self.clear_search).pack(side="left", padx=6)

        columns = ("Code", "Product Name", "Batch Number", "Manufacture Date", "Expiry Date", "Client Name")
        self.search_tree = self._create_tree(self.search_tab, columns)
        self.search_tree.column("Code", width=120)
        self.search_tree.column("Product Name", width=220)
        self.search_tree.column("Batch Number", width=120)
        self.search_tree.column("Manufacture Date", width=130)
        self.search_tree.column("Expiry Date", width=130)
        self.search_tree.column("Client Name", width=180)

    def _build_goods_tab(self):
        top = ttk.Frame(self.goods_tab)
        top.pack(fill="x", pady=(0, 8))
        text_wrap = ttk.Frame(top)
        text_wrap.pack(side="left", fill="x", expand=True)
        ttk.Label(
            text_wrap,
            text="Goods Received Report",
            font=("Helvetica", 12, "bold"),
        ).pack(anchor="w")
        ttk.Label(
            text_wrap,
            text="Code, product name, batch number, manufacture date, expiry date, and client name.",
        ).pack(anchor="w", pady=(2, 0))
        ttk.Button(top, text="Export Professional PDF", command=self.export_goods_received_pdf).pack(side="right", padx=(8, 0))

        columns = ("Code", "Product Name", "Batch Number", "Manufacture Date", "Expiry Date", "Client Name")
        self.goods_tree = self._create_tree(self.goods_tab, columns)
        self.goods_tree.column("Code", width=120)
        self.goods_tree.column("Product Name", width=220)
        self.goods_tree.column("Batch Number", width=120)
        self.goods_tree.column("Manufacture Date", width=130)
        self.goods_tree.column("Expiry Date", width=130)
        self.goods_tree.column("Client Name", width=180)

    def _build_inbound_tab(self):
        top = ttk.Frame(self.inbound_tab)
        top.pack(fill="x", pady=(0, 8))
        text_wrap = ttk.Frame(top)
        text_wrap.pack(side="left", fill="x", expand=True)
        ttk.Label(
            text_wrap,
            text="Inbound Report",
            font=("Helvetica", 12, "bold"),
        ).pack(anchor="w")
        ttk.Label(
            text_wrap,
            text="No., date, client, quantity, type of package stored, and brand.",
        ).pack(anchor="w", pady=(2, 0))
        ttk.Button(top, text="Export Professional PDF", command=self.export_inbound_report_pdf).pack(side="right", padx=(8, 0))

        columns = ("No.", "Date", "Client", "Quantity", "Type of Package Stored", "Brand")
        self.inbound_tree = self._create_tree(self.inbound_tab, columns)
        self.inbound_tree.column("No.", width=60, anchor="center")
        self.inbound_tree.column("Date", width=120)
        self.inbound_tree.column("Client", width=200)
        self.inbound_tree.column("Quantity", width=100, anchor="e")
        self.inbound_tree.column("Type of Package Stored", width=180)
        self.inbound_tree.column("Brand", width=180)

    def _build_outbound_tab(self):
        top = ttk.Frame(self.outbound_tab)
        top.pack(fill="x", pady=(0, 8))

        text_wrap = ttk.Frame(top)
        text_wrap.pack(side="left", fill="x", expand=True)
        ttk.Label(text_wrap, text="Warehouse Outbound Report", font=("Helvetica", 12, "bold")).pack(anchor="w")
        ttk.Label(
            text_wrap,
            text="Warehouse dispatch reporting with batch traceability (no financial data).",
        ).pack(anchor="w", pady=(2, 0))

        btn_wrap = ttk.Frame(top)
        btn_wrap.pack(side="right")
        ttk.Button(btn_wrap, text="Export PDF", command=self.export_outbound_report_pdf).pack(side="left", padx=(0, 6))
        ttk.Button(btn_wrap, text="Export Excel", command=self.export_outbound_report_excel).pack(side="left", padx=(0, 6))
        ttk.Button(btn_wrap, text="Export CSV", command=self.export_outbound_report_csv).pack(side="left", padx=(0, 6))
        ttk.Button(btn_wrap, text="Print Preview", command=self.export_outbound_report_pdf).pack(side="left")

        header = ttk.LabelFrame(self.outbound_tab, text="Report Header", padding=10)
        header.pack(fill="x", pady=(0, 8))
        self.outbound_generated_by = tk.StringVar(value=getpass_user_fallback())
        self.outbound_generated_at = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

        ttk.Label(header, text="Date From").grid(row=0, column=0, sticky="w", padx=(0, 6), pady=4)
        self.outbound_date_from_var = tk.StringVar()
        ttk.Entry(header, textvariable=self.outbound_date_from_var, width=16).grid(row=0, column=1, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Date To").grid(row=0, column=2, sticky="w", padx=(0, 6), pady=4)
        self.outbound_date_to_var = tk.StringVar()
        ttk.Entry(header, textvariable=self.outbound_date_to_var, width=16).grid(row=0, column=3, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Warehouse").grid(row=0, column=4, sticky="w", padx=(0, 6), pady=4)
        self.outbound_warehouse_var = tk.StringVar()
        self.outbound_warehouse_combo = ttk.Combobox(header, textvariable=self.outbound_warehouse_var, values=[""], state="readonly", width=18)
        self.outbound_warehouse_combo.grid(row=0, column=5, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="View").grid(row=0, column=6, sticky="w", padx=(0, 6), pady=4)
        self.outbound_view_var = tk.StringVar(value="Per Shipment")
        self.outbound_view_combo = ttk.Combobox(header, textvariable=self.outbound_view_var, values=["Per Shipment", "Per Destination", "Per Item"], state="readonly", width=16)
        self.outbound_view_combo.grid(row=0, column=7, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Client").grid(row=1, column=0, sticky="w", padx=(0, 6), pady=4)
        self.outbound_client_var = tk.StringVar()
        self.outbound_client_combo = ttk.Combobox(header, textvariable=self.outbound_client_var, values=[""], state="readonly", width=18)
        self.outbound_client_combo.grid(row=1, column=1, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Status").grid(row=1, column=2, sticky="w", padx=(0, 6), pady=4)
        self.outbound_status_var = tk.StringVar()
        self.outbound_status_combo = ttk.Combobox(header, textvariable=self.outbound_status_var, values=[""], state="readonly", width=18)
        self.outbound_status_combo.grid(row=1, column=3, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Destination").grid(row=1, column=4, sticky="w", padx=(0, 6), pady=4)
        self.outbound_destination_var = tk.StringVar()
        self.outbound_destination_combo = ttk.Combobox(header, textvariable=self.outbound_destination_var, values=[""], state="readonly", width=18)
        self.outbound_destination_combo.grid(row=1, column=5, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Search").grid(row=2, column=0, sticky="w", padx=(0, 6), pady=4)
        self.outbound_search_var = tk.StringVar()
        ttk.Entry(header, textvariable=self.outbound_search_var, width=54).grid(row=2, column=1, columnspan=3, sticky="w", padx=(0, 12), pady=4)

        ttk.Label(header, text="Product").grid(row=2, column=4, sticky="w", padx=(0, 6), pady=4)
        self.outbound_product_var = tk.StringVar()
        self.outbound_product_combo = ttk.Combobox(header, textvariable=self.outbound_product_var, values=[""], state="readonly", width=18)
        self.outbound_product_combo.grid(row=2, column=5, sticky="w", padx=(0, 12), pady=4)

        advanced = ttk.LabelFrame(self.outbound_tab, text="Filters", padding=10)
        advanced.pack(fill="x", pady=(0, 8))

        self.outbound_outbound_no_var = tk.StringVar()
        self.outbound_delivery_note_var = tk.StringVar()
        self.outbound_item_code_var = tk.StringVar()
        self.outbound_sku_var = tk.StringVar()
        self.outbound_batch_var = tk.StringVar()
        self.outbound_lot_var = tk.StringVar()
        self.outbound_mfg_var = tk.StringVar()
        self.outbound_exp_var = tk.StringVar()
        self.outbound_requested_by_var = tk.StringVar()
        self.outbound_picked_by_var = tk.StringVar()
        self.outbound_packed_by_var = tk.StringVar()
        self.outbound_dispatched_by_var = tk.StringVar()

        adv_specs = [
            ("Outbound No.", self.outbound_outbound_no_var, 20),
            ("Delivery Note No.", self.outbound_delivery_note_var, 20),
            ("Item Code", self.outbound_item_code_var, 16),
            ("SKU", self.outbound_sku_var, 16),
            ("Batch No.", self.outbound_batch_var, 16),
            ("Lot No.", self.outbound_lot_var, 16),
            ("MFG Date", self.outbound_mfg_var, 16),
            ("EXP Date", self.outbound_exp_var, 16),
            ("Requested By", self.outbound_requested_by_var, 18),
            ("Picked By", self.outbound_picked_by_var, 18),
            ("Packed By", self.outbound_packed_by_var, 18),
            ("Dispatched By", self.outbound_dispatched_by_var, 18),
        ]
        for idx, (label, var, width) in enumerate(adv_specs):
            r = idx // 4
            c = (idx % 4) * 2
            ttk.Label(advanced, text=label).grid(row=r, column=c, sticky="w", padx=(0, 6), pady=4)
            ttk.Entry(advanced, textvariable=var, width=width).grid(row=r, column=c + 1, sticky="w", padx=(0, 12), pady=4)

        row3 = ttk.Frame(header)
        row3.grid(row=3, column=0, columnspan=6, sticky="e", pady=(6, 0))
        ttk.Button(row3, text="Run Report", command=self.run_outbound_report).pack(side="left")
        ttk.Button(row3, text="Clear Filters", command=self.clear_outbound_filters).pack(side="left", padx=6)

        summary = ttk.LabelFrame(self.outbound_tab, text="Summary", padding=10)
        summary.pack(fill="x", pady=(0, 8))
        self.outbound_summary_vars = {
            "total_orders": tk.StringVar(value="0"),
            "total_clients": tk.StringVar(value="0"),
            "total_products": tk.StringVar(value="0"),
            "total_qty": tk.StringVar(value="0"),
            "completed": tk.StringVar(value="0"),
            "pending": tk.StringVar(value="0"),
            "returned": tk.StringVar(value="0"),
            "cancelled": tk.StringVar(value="0"),
        }
        cards = [
            ("Total Outbound Orders", "total_orders"),
            ("Total Clients", "total_clients"),
            ("Total Products Dispatched", "total_products"),
            ("Total Quantity Dispatched", "total_qty"),
            ("Completed Orders", "completed"),
            ("Pending Orders", "pending"),
            ("Returned Orders", "returned"),
            ("Cancelled Orders", "cancelled"),
        ]
        for idx, (label, key) in enumerate(cards):
            card = ttk.Frame(summary, padding=8)
            card.grid(row=idx // 4, column=idx % 4, sticky="ew", padx=6, pady=6)
            ttk.Label(card, text=label).pack(anchor="w")
            ttk.Label(card, textvariable=self.outbound_summary_vars[key], font=("Helvetica", 13, "bold")).pack(anchor="w", pady=(2, 0))
        for col in range(4):
            summary.columnconfigure(col, weight=1)

        table_wrap = ttk.LabelFrame(self.outbound_tab, text="Outbound Summary Table", padding=10)
        table_wrap.pack(fill="both", expand=True, pady=(0, 8))

        self.outbound_summary_tree = ttk.Treeview(table_wrap, columns=(), show=("tree", "headings"), height=12)
        self.outbound_summary_tree.column("#0", width=28, stretch=False)
        self.outbound_summary_tree.heading("#0", text="")
        yscroll = ttk.Scrollbar(table_wrap, orient=tk.VERTICAL, command=self.outbound_summary_tree.yview)
        xscroll = ttk.Scrollbar(table_wrap, orient=tk.HORIZONTAL, command=self.outbound_summary_tree.xview)
        self.outbound_summary_tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.outbound_summary_tree.pack(side="left", fill="both", expand=True)
        yscroll.pack(side="right", fill="y")
        xscroll.pack(side="bottom", fill="x")

        self.outbound_summary_tree.bind("<Double-1>", self._toggle_outbound_expand)
        self.outbound_summary_tree.bind("<<TreeviewSelect>>", self._on_outbound_row_selected)

        pager = ttk.Frame(self.outbound_tab)
        pager.pack(fill="x", pady=(0, 8))
        self.outbound_page_label = tk.StringVar(value="Page 1")
        ttk.Button(pager, text="◀ Prev", command=self._prev_outbound_page).pack(side="left")
        ttk.Button(pager, text="Next ▶", command=self._next_outbound_page).pack(side="left", padx=6)
        ttk.Label(pager, textvariable=self.outbound_page_label).pack(side="left", padx=10)
        ttk.Label(pager, text="Page Size").pack(side="left", padx=(10, 6))
        self.outbound_page_size_var = tk.StringVar(value=str(self.outbound_page_size))
        ttk.Combobox(pager, textvariable=self.outbound_page_size_var, values=["10", "25", "50", "100"], state="readonly", width=6).pack(side="left")
        ttk.Button(pager, text="Apply", command=self._apply_outbound_page_size).pack(side="left", padx=6)

        details_wrap = ttk.LabelFrame(self.outbound_tab, text="Product Details", padding=10)
        details_wrap.pack(fill="both", expand=True, pady=(0, 8))
        detail_cols = (
            "Item Code", "Product Name", "Brand", "SKU", "Batch Number", "Lot Number", "Manufacturing Date", "Expiry Date",
            "Unit / Pack Size", "Quantity Dispatched", "Warehouse", "Bin Location", "Temperature Zone", "Notes"
        )
        self.outbound_details_tree = self._create_tree(details_wrap, detail_cols)
        self.outbound_details_tree.column("Notes", width=300)

        trace_wrap = ttk.LabelFrame(self.outbound_tab, text="Traceability", padding=10)
        trace_wrap.pack(fill="both", expand=True, pady=(0, 8))
        trace_cols = ("Client", "Warehouse", "Product", "Batch Number", "Lot Number", "Manufacturing Date", "Expiry Date", "Quantity", "Dispatch Date", "Destination")
        self.outbound_trace_tree = self._create_tree(trace_wrap, trace_cols)

        dispatch_wrap = ttk.LabelFrame(self.outbound_tab, text="Dispatch Information", padding=10)
        dispatch_wrap.pack(fill="x")
        self.outbound_dispatch_vars = {
            "requested_by": tk.StringVar(value="-"),
            "approved_by": tk.StringVar(value="-"),
            "picked_by": tk.StringVar(value="-"),
            "packed_by": tk.StringVar(value="-"),
            "dispatched_by": tk.StringVar(value="-"),
            "received_by": tk.StringVar(value="-"),
            "dispatch_date": tk.StringVar(value="-"),
            "delivery_date": tk.StringVar(value="-"),
            "delivery_note_no": tk.StringVar(value="-"),
            "goods_transfer_note_no": tk.StringVar(value="-"),
            "vehicle": tk.StringVar(value="-"),
            "courier": tk.StringVar(value="-"),
            "tracking_number": tk.StringVar(value="-"),
        }
        disp_fields = [
            ("Requested By", "requested_by"),
            ("Approved By", "approved_by"),
            ("Picked By", "picked_by"),
            ("Packed By", "packed_by"),
            ("Dispatched By", "dispatched_by"),
            ("Received By", "received_by"),
            ("Dispatch Date", "dispatch_date"),
            ("Delivery Date", "delivery_date"),
            ("Delivery Note No.", "delivery_note_no"),
            ("Goods Transfer Note No.", "goods_transfer_note_no"),
            ("Vehicle", "vehicle"),
            ("Courier", "courier"),
            ("Tracking Number", "tracking_number"),
        ]
        for idx, (label, key) in enumerate(disp_fields):
            r = idx // 4
            c = (idx % 4) * 2
            ttk.Label(dispatch_wrap, text=label).grid(row=r, column=c, sticky="w", padx=(0, 6), pady=3)
            ttk.Label(dispatch_wrap, textvariable=self.outbound_dispatch_vars[key]).grid(row=r, column=c + 1, sticky="w", padx=(0, 12), pady=3)

        doc_wrap = ttk.Frame(dispatch_wrap)
        doc_wrap.grid(row=4, column=0, columnspan=8, sticky="e", pady=(6, 0))
        ttk.Button(doc_wrap, text="Delivery Note", command=self._open_outbound_delivery_note).pack(side="left", padx=(0, 6))
        ttk.Button(doc_wrap, text="Goods Transfer Note", command=self._open_outbound_delivery_note).pack(side="left", padx=(0, 6))
        ttk.Button(doc_wrap, text="Pick List", command=self._open_outbound_pick_list).pack(side="left", padx=(0, 6))
        ttk.Button(doc_wrap, text="Packing List", command=self._open_outbound_packing_list).pack(side="left", padx=(0, 6))
        ttk.Button(doc_wrap, text="Proof of Delivery", command=self._open_outbound_pod).pack(side="left")

        self.outbound_warehouse_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_outbound_choices())
        self.outbound_client_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_outbound_choices())
        self.outbound_status_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_outbound_choices())
        self.outbound_destination_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_outbound_choices())
        self.outbound_product_combo.bind("<<ComboboxSelected>>", lambda _e: self._refresh_outbound_choices())
        self.outbound_view_combo.bind("<<ComboboxSelected>>", lambda _e: self.run_outbound_report())

    def _create_tree(self, parent, columns):
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True)

        tree = ttk.Treeview(frame, columns=columns, show="headings", height=18)
        for col in columns:
            tree.heading(col, text=col)
            tree.column(col, width=120, stretch=True)

        yscroll = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=tree.yview)
        xscroll = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)

        tree.pack(side="left", fill="both", expand=True)
        yscroll.pack(side="right", fill="y")
        xscroll.pack(side="bottom", fill="x")
        return tree

    def refresh_all(self):
        self.inventory_manager = InventoryManager(self.data_folder)
        self.goods_rows = self._build_goods_rows()
        self.inbound_rows = self._build_inbound_rows()
        self.outbound_rows = self._load_outbound_rows()
        self.run_search()
        self._populate_tree(self.goods_tree, self.goods_rows, ("code", "product_name", "batch_number", "manufacture_date", "expiry_date", "client_name"))
        self._populate_tree(self.inbound_tree, self.inbound_rows, ("row_no", "date", "client", "quantity", "package_type", "brand"))
        self._refresh_outbound_choices()
        self.run_outbound_report()
        self.status_var.set(
            f"Inventory Search: {len(self.search_rows)} | Goods Received: {len(self.goods_rows)} | Inbound Report: {len(self.inbound_rows)} | Outbound Report: {len(self.outbound_rows)}"
        )

    def run_search(self):
        code_q = self._norm(self.code_var.get())
        product_q = self._norm(self.product_var.get())
        batch_q = self._norm(self.batch_var.get())
        manufacture_q = self._norm(self.manufacture_var.get())
        expiry_q = self._norm(self.expiry_var.get())
        client_q = self._norm(self.client_var.get())

        rows = []
        for item in self.inventory_manager.items:
            row = self._item_to_goods_row(item)
            if code_q and code_q not in self._norm(row["code"]):
                continue
            if product_q and product_q not in self._norm(row["product_name"]):
                continue
            if batch_q and batch_q not in self._norm(row["batch_number"]):
                continue
            if manufacture_q and manufacture_q not in self._norm(row["manufacture_date"]):
                continue
            if expiry_q and expiry_q not in self._norm(row["expiry_date"]):
                continue
            if client_q and client_q not in self._norm(row["client_name"]):
                continue
            rows.append(row)

        self.search_rows = rows
        self._populate_tree(
            self.search_tree,
            rows,
            ("code", "product_name", "batch_number", "manufacture_date", "expiry_date", "client_name"),
        )
        self.status_var.set(
            f"Inventory Search: {len(self.search_rows)} | Goods Received: {len(self.goods_rows)} | Inbound Report: {len(self.inbound_rows)}"
        )

    def clear_search(self):
        self.code_var.set("")
        self.product_var.set("")
        self.batch_var.set("")
        self.manufacture_var.set("")
        self.expiry_var.set("")
        self.client_var.set("")
        self.run_search()

    def _populate_tree(self, tree, rows, keys):
        for child in tree.get_children():
            tree.delete(child)
        for row in rows:
            tree.insert("", "end", values=tuple(row.get(key, "") for key in keys))

    def _build_goods_rows(self):
        rows = [self._item_to_goods_row(item) for item in self.inventory_manager.items]
        rows.sort(key=lambda row: self._sort_key(row.get("received_date")), reverse=True)
        return rows

    def _build_inbound_rows(self):
        rows = []
        for index, item in enumerate(sorted(self.inventory_manager.items, key=lambda itm: self._sort_key(self._item_date(itm)), reverse=True), start=1):
            package_type = self._clean_text(getattr(item, "package_type", ""))
            if not package_type:
                package_type = self._infer_package_type(item)
            rows.append({
                "row_no": index,
                "date": self._item_date(item),
                "client": self._item_client(item),
                "quantity": self._format_quantity(getattr(item, "quantity", 0)),
                "package_type": package_type or "Others",
                "brand": self._clean_text(getattr(item, "brand", "")) or "-",
            })
        return rows

    def _item_to_goods_row(self, item):
        return {
            "code": self._clean_text(getattr(item, "item_id", "")),
            "product_name": self._clean_text(getattr(item, "name", "")),
            "batch_number": self._clean_text(getattr(item, "batch_number", "")),
            "manufacture_date": self._clean_text(getattr(item, "manufacture_date", "")),
            "expiry_date": self._clean_text(getattr(item, "expiry_date", "")),
            "client_name": self._item_client(item),
            "received_date": self._item_date(item),
        }

    def _item_client(self, item):
        return self._clean_text(getattr(item, "client_name", "")) or self._clean_text(getattr(item, "supplier", "")) or "-"

    def _item_date(self, item):
        return self._clean_text(getattr(item, "invoice_date", "")) or self._clean_text(getattr(item, "created_date", "")) or ""

    def _infer_package_type(self, item):
        text = " ".join(
            [
                self._clean_text(getattr(item, "name", "")),
                self._clean_text(getattr(item, "description", "")),
                self._clean_text(getattr(item, "category", "")),
            ]
        ).lower()
        if "carton" in text or "cartoon" in text:
            return "Carton"
        if "box" in text:
            return "Box"
        if "pack" in text:
            return "Pack"
        if "bottle" in text:
            return "Bottle"
        if "tube" in text:
            return "Tube"
        return "Others"

    def _format_quantity(self, value):
        try:
            value = float(value or 0)
            if value.is_integer():
                return int(value)
            return f"{value:.2f}"
        except Exception:
            return value

    def _sort_key(self, value):
        raw = self._clean_text(value)
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(raw, fmt)
            except Exception:
                pass
        return datetime.min

    def _clean_text(self, value):
        return str(value or "").strip()

    def _norm(self, value):
        return self._clean_text(value).lower()

    def export_goods_received_pdf(self):
        rows = list(self.goods_rows or [])
        if not rows:
            messagebox.showwarning("Export PDF", "No goods received rows to export.")
            return
        unique_clients = len({self._clean_text(row.get("client_name")) for row in rows if self._clean_text(row.get("client_name"))})
        unique_batches = len({self._clean_text(row.get("batch_number")) for row in rows if self._clean_text(row.get("batch_number"))})
        path = self._build_report_pdf(
            report_name="Goods_Received_Report",
            title="Goods Received Report",
            subtitle="Inventory intake register with manufacturing, expiry, and client traceability.",
            headers=["Code", "Product Name", "Batch Number", "Manufacture Date", "Expiry Date", "Client Name"],
            rows=[[row.get("code", ""), row.get("product_name", ""), row.get("batch_number", ""), row.get("manufacture_date", ""), row.get("expiry_date", ""), row.get("client_name", "")] for row in rows],
            col_widths=[70, 185, 90, 95, 95, 125],
            summary_pairs=[
                ("Total Rows", str(len(rows))),
                ("Unique Clients", str(unique_clients)),
                ("Unique Batches", str(unique_batches)),
                ("Generated By", getpass_user_fallback()),
            ],
        )
        if path:
            self._open_exported_pdf(path)

    def export_inbound_report_pdf(self):
        rows = list(self.inbound_rows or [])
        if not rows:
            messagebox.showwarning("Export PDF", "No inbound rows to export.")
            return
        total_quantity = 0.0
        for row in rows:
            try:
                total_quantity += float(row.get("quantity") or 0)
            except Exception:
                pass
        unique_clients = len({self._clean_text(row.get("client")) for row in rows if self._clean_text(row.get("client"))})
        unique_brands = len({self._clean_text(row.get("brand")) for row in rows if self._clean_text(row.get("brand")) and self._clean_text(row.get("brand")) != "-"})
        path = self._build_report_pdf(
            report_name="Inbound_Report",
            title="Inbound Report",
            subtitle="Storage intake summary showing client, quantity, package type, and brand.",
            headers=["No.", "Date", "Client", "Quantity", "Type of Package Stored", "Brand"],
            rows=[[row.get("row_no", ""), row.get("date", ""), row.get("client", ""), row.get("quantity", ""), row.get("package_type", ""), row.get("brand", "")] for row in rows],
            col_widths=[40, 85, 165, 75, 180, 125],
            summary_pairs=[
                ("Total Rows", str(len(rows))),
                ("Total Quantity", self._format_quantity(total_quantity)),
                ("Unique Clients", str(unique_clients)),
                ("Unique Brands", str(unique_brands)),
            ],
        )
        if path:
            self._open_exported_pdf(path)

    def _build_report_pdf(self, report_name, title, subtitle, headers, rows, col_widths, summary_pairs=None):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Export PDF", "PDF export is not available in this build.")
            return None
        try:
            reports_folder = os.path.join(self.data_folder, "InventoryReports")
            os.makedirs(reports_folder, exist_ok=True)
            filepath = os.path.join(reports_folder, f"{report_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "InventoryReportTitle",
                parent=styles["Heading1"],
                fontSize=16,
                alignment=1,
                textColor=colors.HexColor("#2c3e50"),
            )
            subtitle_style = ParagraphStyle(
                "InventoryReportSub",
                parent=styles["Normal"],
                fontSize=9,
                alignment=1,
                textColor=colors.grey,
            )
            info_style = ParagraphStyle(
                "InventoryReportInfo",
                parent=styles["Normal"],
                fontSize=8.5,
                alignment=1,
                textColor=colors.HexColor("#566573"),
            )
            cell_style = ParagraphStyle(
                "InventoryReportCell",
                parent=styles["BodyText"],
                fontSize=8,
                leading=9,
                wordWrap="CJK",
            )
            summary_label_style = ParagraphStyle(
                "InventorySummaryLabel",
                parent=styles["BodyText"],
                fontSize=8,
                textColor=colors.HexColor("#5d6d7e"),
            )
            summary_value_style = ParagraphStyle(
                "InventorySummaryValue",
                parent=styles["BodyText"],
                fontSize=12,
                textColor=colors.HexColor("#0b5394"),
            )

            doc = SimpleDocTemplate(
                filepath,
                pagesize=landscape(A4),
                leftMargin=0.35 * inch,
                rightMargin=0.35 * inch,
                topMargin=0.4 * inch,
                bottomMargin=0.4 * inch,
            )

            elements = []
            elements.extend(build_pdf_logo_flowables(self.data_folder, width=1.0 * inch, height=1.0 * inch, spacer_height=0.08 * inch))
            elements.extend([
                Paragraph(title, title_style),
                Paragraph(subtitle, subtitle_style),
                Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Hope Pharma Inventory Reports", info_style),
                Spacer(1, 8),
            ])

            summary_pairs = list(summary_pairs or [])
            if summary_pairs:
                summary_cells = []
                for label, value in summary_pairs:
                    summary_cells.append(
                        [
                            Paragraph(str(label), summary_label_style),
                            Paragraph(f"<b>{self._clean_text(value) or '-'}</b>", summary_value_style),
                        ]
                    )
                if len(summary_cells) % 2 == 1:
                    summary_cells.append(
                        [
                            Paragraph("", summary_label_style),
                            Paragraph("", summary_value_style),
                        ]
                    )
                summary_rows = []
                for i in range(0, len(summary_cells), 2):
                    left = summary_cells[i]
                    right = summary_cells[i + 1]
                    summary_rows.append([left[0], left[1], right[0], right[1]])

                summary_table = Table(summary_rows, colWidths=[doc.width * 0.17, doc.width * 0.18, doc.width * 0.17, doc.width * 0.18])
                summary_table.setStyle(TableStyle([
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d5d8dc")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e5e7e9")),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8f9fb")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]))
                elements.append(summary_table)
                elements.append(Spacer(1, 10))

            table_data = [headers]
            for row in rows:
                table_data.append([self._pdf_paragraph(value, cell_style) for value in row])

            table = Table(table_data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b5394")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 8))
            elements.append(
                Paragraph(
                    "This report is generated by Hope Pharma and includes official branding for inventory traceability and audit use.",
                    info_style,
                )
            )
            doc.build(elements)
            return filepath
        except Exception as exc:
            messagebox.showerror("Export PDF", f"Failed to generate PDF:\n{exc}")
            return None

    def _pdf_paragraph(self, value, style):
        text = self._clean_text(value) or "-"
        safe_text = (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        return Paragraph(safe_text, style)

    def _open_exported_pdf(self, path):
        try:
            webbrowser.open("file://" + os.path.abspath(path))
        except Exception:
            pass
        messagebox.showinfo("Export PDF", f"PDF saved:\n{path}")

    def _load_outbound_rows(self):
        def _read_list(path):
            try:
                if os.path.exists(path):
                    with open(path, "r", encoding="utf-8") as handle:
                        payload = json.load(handle)
                    if isinstance(payload, list):
                        return payload
            except Exception:
                pass
            return []

        def _normalize_item_from_inventory(item, quantity):
            return {
                "item_code": self._clean_text(getattr(item, "item_id", "")),
                "product_name": self._clean_text(getattr(item, "name", "")),
                "brand": self._clean_text(getattr(item, "brand", "")),
                "sku": self._clean_text(getattr(item, "sku", "")),
                "batch_number": self._clean_text(getattr(item, "batch_number", "")),
                "lot_number": self._clean_text(getattr(item, "lot_number", "")),
                "manufacture_date": self._clean_text(getattr(item, "manufacture_date", "")),
                "expiry_date": self._clean_text(getattr(item, "expiry_date", "")),
                "unit_pack_size": self._clean_text(" / ".join([x for x in [getattr(item, "unit_of_measure", ""), getattr(item, "pack_size", "")] if self._clean_text(x)])),
                "warehouse": self._clean_text(getattr(item, "warehouse_code", "")),
                "bin_location": self._clean_text(getattr(item, "bin_location", "")),
                "temperature_zone": "Temperature Controlled" if bool(getattr(item, "temperature_controlled", False)) else "-",
                "notes": "",
                "quantity_dispatched": self._clean_text(quantity),
                "raw_quantity_dispatched": quantity,
            }

        rows = []

        for raw in _read_list(os.path.join(self.data_folder, "outbounds.json")):
            outbound_no = self._clean_text(raw.get("outbound_no"))
            if not outbound_no:
                continue
            items = list(raw.get("items") or [])
            normalized_items = []
            for line in items:
                normalized_items.append({
                    "item_code": self._clean_text(line.get("item_id") or line.get("source_item_id")),
                    "product_name": self._clean_text(line.get("product_name")),
                    "brand": self._clean_text(line.get("brand")),
                    "sku": self._clean_text(line.get("sku")),
                    "batch_number": self._clean_text(line.get("batch_number")),
                    "lot_number": self._clean_text(line.get("lot_number")),
                    "manufacture_date": self._clean_text(line.get("manufacture_date")),
                    "expiry_date": self._clean_text(line.get("expiry_date")),
                    "unit_pack_size": self._clean_text(line.get("unit_pack_size")) or self._clean_text(" / ".join([x for x in [line.get("unit_of_measure"), line.get("pack_size")] if self._clean_text(x)])),
                    "warehouse": self._clean_text(line.get("warehouse_code") or raw.get("warehouse_code")),
                    "bin_location": self._clean_text(line.get("bin_location") or raw.get("bin_location")),
                    "temperature_zone": "Temperature Controlled" if bool(line.get("temperature_controlled")) else "-",
                    "notes": self._clean_text(raw.get("notes")),
                    "quantity_dispatched": self._clean_text(line.get("dispatched_quantity") or line.get("packed_quantity") or line.get("picked_quantity") or line.get("requested_quantity")),
                    "raw_quantity_dispatched": line.get("dispatched_quantity") or line.get("packed_quantity") or line.get("picked_quantity") or line.get("requested_quantity") or 0,
                })
            rows.append({
                "outbound_no": outbound_no,
                "delivery_note_no": self._clean_text(raw.get("delivery_note_no")),
                "date": self._clean_text(raw.get("date")),
                "client": self._clean_text(raw.get("client_owner")),
                "destination": self._clean_text(raw.get("destination")),
                "warehouse": self._clean_text(raw.get("warehouse_code")),
                "bin_location": self._clean_text(raw.get("bin_location")),
                "status": self._clean_text(raw.get("status")),
                "requested_by": self._clean_text(raw.get("requested_by")),
                "approved_by": self._clean_text(raw.get("approved_by")),
                "picked_by": self._clean_text(raw.get("picked_by")),
                "packed_by": self._clean_text(raw.get("packed_by")),
                "dispatched_by": self._clean_text(raw.get("dispatched_by")),
                "received_by": self._clean_text(raw.get("received_by")),
                "dispatch_date": self._clean_text(raw.get("dispatched_date")),
                "delivery_date": self._clean_text(raw.get("received_date")),
                "notes": self._clean_text(raw.get("notes")),
                "vehicle": self._clean_text(raw.get("vehicle")),
                "courier": self._clean_text(raw.get("courier")),
                "tracking_number": self._clean_text(raw.get("tracking_number")),
                "source": "outbound",
                "items": normalized_items,
            })

        transfer_status_map = {
            "Requested": "Draft",
            "Pending": "Pending Approval",
            "Approved": "Approved",
            "Dispatched": "Dispatched",
            "Received": "Delivered",
        }
        for trf in _read_list(os.path.join(self.data_folder, "stock_transfers.json")):
            if self._clean_text(trf.get("transfer_type")) != "Client Transfer":
                continue
            transfer_id = self._clean_text(trf.get("transfer_id"))
            if not transfer_id:
                continue
            status = transfer_status_map.get(self._clean_text(trf.get("status")), self._clean_text(trf.get("status")) or "Draft")
            qty = trf.get("quantity") or 0
            items = [{
                "item_code": self._clean_text(trf.get("item_id")),
                "product_name": self._clean_text(trf.get("product_name")),
                "brand": self._clean_text(trf.get("brand")),
                "sku": self._clean_text(trf.get("sku")),
                "batch_number": self._clean_text(trf.get("batch_number")),
                "lot_number": self._clean_text(trf.get("lot_number")),
                "manufacture_date": self._clean_text(trf.get("manufacture_date")),
                "expiry_date": self._clean_text(trf.get("expiry_date")),
                "unit_pack_size": self._clean_text(" / ".join([x for x in [trf.get("unit_of_measure"), trf.get("pack_size")] if self._clean_text(x)])),
                "warehouse": self._clean_text(trf.get("warehouse_from")),
                "bin_location": self._clean_text(trf.get("bin_from")),
                "temperature_zone": "-",
                "notes": self._clean_text(trf.get("transfer_description") or trf.get("notes")),
                "quantity_dispatched": self._clean_text(qty),
                "raw_quantity_dispatched": qty,
            }]
            rows.append({
                "outbound_no": transfer_id,
                "delivery_note_no": transfer_id,
                "date": self._clean_text(trf.get("requested_date")),
                "client": self._clean_text(trf.get("stored_for_client")),
                "destination": self._clean_text(trf.get("destination_client")),
                "warehouse": self._clean_text(trf.get("warehouse_from")),
                "bin_location": self._clean_text(trf.get("bin_from")),
                "status": status,
                "requested_by": self._clean_text(trf.get("requested_by")),
                "approved_by": self._clean_text(trf.get("approved_by")),
                "picked_by": "",
                "packed_by": "",
                "dispatched_by": self._clean_text(trf.get("approved_by")),
                "received_by": self._clean_text(trf.get("received_by")),
                "dispatch_date": self._clean_text(trf.get("approved_date") or trf.get("requested_date")),
                "delivery_date": self._clean_text(trf.get("received_date")),
                "notes": self._clean_text(trf.get("notes")),
                "vehicle": "",
                "courier": "",
                "tracking_number": "",
                "source": "goods_transfer",
                "items": items,
            })

        sales = _read_list(os.path.join(self.data_folder, "sales_records.json"))
        sales_groups = {}
        for s in sales:
            inv_id = self._clean_text(s.get("invoice_id"))
            sale_id = self._clean_text(s.get("sale_id"))
            key = inv_id or sale_id
            if not key:
                continue
            sales_groups.setdefault(key, []).append(s)
        for key, group in sales_groups.items():
            customer = self._clean_text((group[0] or {}).get("customer_name"))
            sale_date = self._clean_text((group[0] or {}).get("sale_date"))
            items = []
            for s in group:
                item_id = self._clean_text(s.get("item_id"))
                qty = 0.0
                try:
                    qty = float(s.get("quantity") or 0)
                except Exception:
                    qty = 0.0
                try:
                    qty += float(s.get("promotion_quantity") or 0)
                except Exception:
                    pass
                inv_item = self.inventory_manager.get_item(item_id) if hasattr(self.inventory_manager, "get_item") else None
                if inv_item:
                    items.append(_normalize_item_from_inventory(inv_item, qty))
                else:
                    items.append({
                        "item_code": item_id,
                        "product_name": self._clean_text(s.get("item_name")),
                        "brand": "",
                        "sku": "",
                        "batch_number": "",
                        "lot_number": "",
                        "manufacture_date": "",
                        "expiry_date": "",
                        "unit_pack_size": "",
                        "warehouse": "",
                        "bin_location": "",
                        "temperature_zone": "-",
                        "notes": "",
                        "quantity_dispatched": self._clean_text(qty),
                        "raw_quantity_dispatched": qty,
                    })
            outbound_no = f"SAL-{key}"
            rows.append({
                "outbound_no": outbound_no,
                "delivery_note_no": key,
                "date": sale_date,
                "client": customer,
                "destination": customer,
                "warehouse": "",
                "bin_location": "",
                "status": "Delivered",
                "requested_by": "Sales",
                "approved_by": "Sales",
                "picked_by": "",
                "packed_by": "",
                "dispatched_by": "Sales",
                "received_by": customer,
                "dispatch_date": sale_date,
                "delivery_date": sale_date,
                "notes": "",
                "vehicle": "",
                "courier": "",
                "tracking_number": "",
                "source": "sales",
                "items": items,
            })

        rows.sort(key=lambda r: self._sort_key(r.get("date")), reverse=True)
        return rows

    def clear_outbound_filters(self):
        self.outbound_date_from_var.set("")
        self.outbound_date_to_var.set("")
        self.outbound_warehouse_var.set("")
        self.outbound_client_var.set("")
        self.outbound_status_var.set("")
        self.outbound_destination_var.set("")
        self.outbound_product_var.set("")
        self.outbound_search_var.set("")
        self.outbound_outbound_no_var.set("")
        self.outbound_delivery_note_var.set("")
        self.outbound_item_code_var.set("")
        self.outbound_sku_var.set("")
        self.outbound_batch_var.set("")
        self.outbound_lot_var.set("")
        self.outbound_mfg_var.set("")
        self.outbound_exp_var.set("")
        self.outbound_requested_by_var.set("")
        self.outbound_picked_by_var.set("")
        self.outbound_packed_by_var.set("")
        self.outbound_dispatched_by_var.set("")
        self.outbound_page = 1
        self._refresh_outbound_choices()
        self.run_outbound_report()

    def _outbound_current_filters(self):
        return {
            "date_from": self._clean_text(self.outbound_date_from_var.get()),
            "date_to": self._clean_text(self.outbound_date_to_var.get()),
            "warehouse": self._clean_text(self.outbound_warehouse_var.get()),
            "client": self._clean_text(self.outbound_client_var.get()),
            "status": self._clean_text(self.outbound_status_var.get()),
            "destination": self._clean_text(self.outbound_destination_var.get()),
            "product": self._clean_text(self.outbound_product_var.get()),
            "search": self._clean_text(self.outbound_search_var.get()),
            "outbound_no": self._clean_text(self.outbound_outbound_no_var.get()),
            "delivery_note_no": self._clean_text(self.outbound_delivery_note_var.get()),
            "item_code": self._clean_text(self.outbound_item_code_var.get()),
            "sku": self._clean_text(self.outbound_sku_var.get()),
            "batch_number": self._clean_text(self.outbound_batch_var.get()),
            "lot_number": self._clean_text(self.outbound_lot_var.get()),
            "manufacture_date": self._clean_text(self.outbound_mfg_var.get()),
            "expiry_date": self._clean_text(self.outbound_exp_var.get()),
            "requested_by": self._clean_text(self.outbound_requested_by_var.get()),
            "picked_by": self._clean_text(self.outbound_picked_by_var.get()),
            "packed_by": self._clean_text(self.outbound_packed_by_var.get()),
            "dispatched_by": self._clean_text(self.outbound_dispatched_by_var.get()),
        }

    def _refresh_outbound_choices(self):
        filters = self._outbound_current_filters() if hasattr(self, "outbound_date_from_var") else {}
        wh_probe = self._norm(filters.get("warehouse"))
        client_probe = self._norm(filters.get("client"))
        status_probe = self._norm(filters.get("status"))
        dest_probe = self._norm(filters.get("destination"))
        product_probe = self._norm(filters.get("product"))

        warehouses = set()
        clients = set()
        statuses = set()
        destinations = set()
        products = set()

        for row in self.outbound_rows:
            if wh_probe and wh_probe != self._norm(row.get("warehouse")):
                continue
            if client_probe and client_probe != self._norm(row.get("client")):
                continue
            if status_probe and status_probe != self._norm(row.get("status")):
                continue
            if dest_probe and dest_probe != self._norm(row.get("destination")):
                continue
            if product_probe:
                if not any(product_probe == self._norm(it.get("product_name")) for it in (row.get("items") or [])):
                    continue
            warehouses.add(self._clean_text(row.get("warehouse")))
            clients.add(self._clean_text(row.get("client")))
            statuses.add(self._clean_text(row.get("status")))
            destinations.add(self._clean_text(row.get("destination")))
            for it in row.get("items") or []:
                products.add(self._clean_text(it.get("product_name")))

        self.outbound_warehouse_combo.configure(values=[""] + sorted({x for x in warehouses if x}))
        self.outbound_client_combo.configure(values=[""] + sorted({x for x in clients if x}))
        self.outbound_status_combo.configure(values=[""] + sorted({x for x in statuses if x}))
        self.outbound_destination_combo.configure(values=[""] + sorted({x for x in destinations if x}))
        self.outbound_product_combo.configure(values=[""] + sorted({x for x in products if x}))

        for key, combo in (
            ("warehouse", self.outbound_warehouse_combo),
            ("client", self.outbound_client_combo),
            ("status", self.outbound_status_combo),
            ("destination", self.outbound_destination_combo),
            ("product", self.outbound_product_combo),
        ):
            current = self._clean_text(filters.get(key))
            values = list(combo.cget("values") or [])
            if current and current not in values:
                combo.set("")

    def _parse_date_filter(self, value):
        return self._sort_key(value) if self._clean_text(value) else None

    def _outbound_matches_search(self, outbound, query):
        q = self._norm(query)
        if not q:
            return True
        hay = " ".join([
            self._clean_text(outbound.get("outbound_no")),
            self._clean_text(outbound.get("delivery_note_no")),
            self._clean_text(outbound.get("client")),
            self._clean_text(outbound.get("destination")),
            self._clean_text(outbound.get("warehouse")),
            self._clean_text(outbound.get("status")),
            self._clean_text(outbound.get("requested_by")),
            self._clean_text(outbound.get("picked_by")),
            self._clean_text(outbound.get("packed_by")),
            self._clean_text(outbound.get("dispatched_by")),
            self._clean_text(outbound.get("received_by")),
        ]).lower()
        if q in hay:
            return True
        for it in outbound.get("items") or []:
            item_hay = " ".join([
                self._clean_text(it.get("item_code")),
                self._clean_text(it.get("product_name")),
                self._clean_text(it.get("brand")),
                self._clean_text(it.get("sku")),
                self._clean_text(it.get("batch_number")),
                self._clean_text(it.get("lot_number")),
            ]).lower()
            if q in item_hay:
                return True
        return False

    def run_outbound_report(self):
        self.outbound_generated_at.set(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        filters = self._outbound_current_filters()
        date_from = self._parse_date_filter(filters.get("date_from"))
        date_to = self._parse_date_filter(filters.get("date_to"))
        wh = self._norm(filters.get("warehouse"))
        client = self._norm(filters.get("client"))
        status = self._norm(filters.get("status"))
        destination = self._norm(filters.get("destination"))
        product = self._norm(filters.get("product"))
        outbound_no_q = self._norm(filters.get("outbound_no"))
        delivery_note_q = self._norm(filters.get("delivery_note_no"))
        requested_by_q = self._norm(filters.get("requested_by"))
        picked_by_q = self._norm(filters.get("picked_by"))
        packed_by_q = self._norm(filters.get("packed_by"))
        dispatched_by_q = self._norm(filters.get("dispatched_by"))

        item_code_q = self._norm(filters.get("item_code"))
        sku_q = self._norm(filters.get("sku"))
        batch_q = self._norm(filters.get("batch_number"))
        lot_q = self._norm(filters.get("lot_number"))
        mfg_q = self._norm(filters.get("manufacture_date"))
        exp_q = self._norm(filters.get("expiry_date"))

        out = []
        for row in self.outbound_rows:
            row_date = self._sort_key(row.get("date"))
            if date_from and row_date < date_from:
                continue
            if date_to and row_date > date_to:
                continue
            if wh and wh != self._norm(row.get("warehouse")):
                continue
            if client and client != self._norm(row.get("client")):
                continue
            if status and status != self._norm(row.get("status")):
                continue
            if destination and destination != self._norm(row.get("destination")):
                continue
            if outbound_no_q and outbound_no_q not in self._norm(row.get("outbound_no")):
                continue
            if delivery_note_q and delivery_note_q not in self._norm(row.get("delivery_note_no")):
                continue
            if requested_by_q and requested_by_q not in self._norm(row.get("requested_by")):
                continue
            if picked_by_q and picked_by_q not in self._norm(row.get("picked_by")):
                continue
            if packed_by_q and packed_by_q not in self._norm(row.get("packed_by")):
                continue
            if dispatched_by_q and dispatched_by_q not in self._norm(row.get("dispatched_by")):
                continue
            if product and not any(product == self._norm(it.get("product_name")) for it in (row.get("items") or [])):
                continue
            if any([item_code_q, sku_q, batch_q, lot_q, mfg_q, exp_q]):
                matched = False
                for it in row.get("items") or []:
                    if item_code_q and item_code_q not in self._norm(it.get("item_code")):
                        continue
                    if sku_q and sku_q not in self._norm(it.get("sku")):
                        continue
                    if batch_q and batch_q not in self._norm(it.get("batch_number")):
                        continue
                    if lot_q and lot_q not in self._norm(it.get("lot_number")):
                        continue
                    if mfg_q and mfg_q not in self._norm(it.get("manufacture_date")):
                        continue
                    if exp_q and exp_q not in self._norm(it.get("expiry_date")):
                        continue
                    matched = True
                    break
                if not matched:
                    continue
            if not self._outbound_matches_search(row, filters.get("search")):
                continue
            out.append(row)

        self.outbound_filtered = out
        self._build_outbound_view_rows()
        self._apply_outbound_summary()
        self._populate_outbound_summary_tree()
        self._populate_outbound_traceability()

    def _outbound_view_mode(self):
        return self._clean_text(getattr(self, "outbound_view_var", tk.StringVar(value="Per Shipment")).get() if hasattr(self, "outbound_view_var") else "Per Shipment") or "Per Shipment"

    def _build_outbound_view_rows(self):
        mode = self._outbound_view_mode()
        self._outbound_view_rows = []
        self._outbound_view_map = {}

        if mode == "Per Shipment":
            for ob in self.outbound_filtered:
                items = ob.get("items") or []
                qty = 0.0
                for it in items:
                    try:
                        qty += float(it.get("raw_quantity_dispatched") or 0)
                    except Exception:
                        pass
                view_row = {
                    "key": self._clean_text(ob.get("outbound_no")),
                    "outbound_no": ob.get("outbound_no"),
                    "date": ob.get("date"),
                    "client": ob.get("client"),
                    "destination": ob.get("destination"),
                    "warehouse": ob.get("warehouse"),
                    "num_products": len(items),
                    "total_qty": qty,
                    "status": ob.get("status"),
                    "_shipments": [ob],
                    "_items": items,
                }
                self._outbound_view_rows.append(view_row)
            return

        if mode == "Per Item":
            for ob in self.outbound_filtered:
                for it in ob.get("items") or []:
                    qty = 0.0
                    try:
                        qty = float(it.get("raw_quantity_dispatched") or 0)
                    except Exception:
                        qty = 0.0
                    key = f"{self._clean_text(ob.get('outbound_no'))}::{self._clean_text(it.get('item_code'))}::{self._clean_text(it.get('batch_number'))}"
                    self._outbound_view_rows.append({
                        "key": key,
                        "outbound_no": ob.get("outbound_no"),
                        "date": ob.get("date"),
                        "client": ob.get("client"),
                        "destination": ob.get("destination"),
                        "warehouse": ob.get("warehouse"),
                        "item_code": it.get("item_code"),
                        "product_name": it.get("product_name"),
                        "batch_number": it.get("batch_number"),
                        "qty": qty,
                        "status": ob.get("status"),
                        "_shipments": [ob],
                        "_items": [it],
                    })
            return

        status_priority = [
            "Returned",
            "Cancelled",
            "Delivered",
            "Dispatched",
            "Ready for Dispatch",
            "Packed",
            "Packing",
            "Picked",
            "Picking",
            "Reserved",
            "Approved",
            "Pending Approval",
            "Draft",
        ]
        priority_map = {self._norm(x): i for i, x in enumerate(status_priority)}

        groups = {}
        for ob in self.outbound_filtered:
            dest = self._clean_text(ob.get("destination"))
            client = self._clean_text(ob.get("client"))
            warehouse = self._clean_text(ob.get("warehouse"))
            key = (self._norm(dest), self._norm(client), self._norm(warehouse))
            grp = groups.get(key)
            if not grp:
                grp = {
                    "destination": dest,
                    "client": client,
                    "warehouse": warehouse,
                    "shipments": [],
                    "items": [],
                    "statuses": [],
                    "dates": [],
                }
                groups[key] = grp
            grp["shipments"].append(ob)
            grp["statuses"].append(self._clean_text(ob.get("status")))
            grp["dates"].append(self._clean_text(ob.get("date")))
            for it in ob.get("items") or []:
                grp["items"].append(it)

        for key, grp in groups.items():
            qty = 0.0
            for it in grp["items"]:
                try:
                    qty += float(it.get("raw_quantity_dispatched") or 0)
                except Exception:
                    pass
            status = ""
            best = None
            for st in grp["statuses"]:
                idx = priority_map.get(self._norm(st))
                if idx is None:
                    continue
                if best is None or idx < best:
                    best = idx
                    status = st
            dates = [d for d in grp["dates"] if d]
            date_display = ""
            if dates:
                dates_sorted = sorted(dates)
                date_display = dates_sorted[-1]
            self._outbound_view_rows.append({
                "key": f"{grp['destination']}::{grp['client']}::{grp['warehouse']}",
                "destination": grp["destination"],
                "client": grp["client"],
                "warehouse": grp["warehouse"],
                "date": date_display,
                "shipments": len(grp["shipments"]),
                "num_products": len(grp["items"]),
                "total_qty": qty,
                "status": status or "Draft",
                "_shipments": grp["shipments"],
                "_items": grp["items"],
            })

    def _apply_outbound_summary(self):
        rows = self.outbound_filtered
        total_orders = len(rows)
        total_clients = len({self._clean_text(r.get("client")) for r in rows if self._clean_text(r.get("client"))})
        total_products = 0
        total_qty = 0.0
        completed = 0
        returned = 0
        cancelled = 0
        pending = 0

        for r in rows:
            status = self._clean_text(r.get("status"))
            if status == "Delivered":
                completed += 1
            elif status == "Returned":
                returned += 1
            elif status == "Cancelled":
                cancelled += 1
            else:
                pending += 1

            items = r.get("items") or []
            total_products += len(items)
            for it in items:
                if self._clean_text(r.get("status")) in {"Dispatched", "Delivered", "Returned"}:
                    try:
                        total_qty += float(it.get("raw_quantity_dispatched") or 0)
                    except Exception:
                        pass

        mode = self._outbound_view_mode()
        total_label = "total_orders"
        if mode == "Per Destination":
            self.outbound_summary_vars["total_orders"].set(str(len(self._outbound_view_rows)))
        else:
            self.outbound_summary_vars["total_orders"].set(str(total_orders))
        self.outbound_summary_vars["total_clients"].set(str(total_clients))
        self.outbound_summary_vars["total_products"].set(str(total_products))
        self.outbound_summary_vars["total_qty"].set(self._format_quantity(total_qty))
        self.outbound_summary_vars["completed"].set(str(completed))
        self.outbound_summary_vars["pending"].set(str(pending))
        self.outbound_summary_vars["returned"].set(str(returned))
        self.outbound_summary_vars["cancelled"].set(str(cancelled))

    def _ensure_outbound_status_tags(self):
        if getattr(self, "_outbound_tags_ready", False):
            return
        palette = {
            "Draft": ("#f3f4f6", "#111827"),
            "Pending Approval": ("#fff7ed", "#9a3412"),
            "Approved": ("#eff6ff", "#1d4ed8"),
            "Reserved": ("#eef2ff", "#4338ca"),
            "Picking": ("#ecfeff", "#155e75"),
            "Picked": ("#ecfeff", "#155e75"),
            "Packing": ("#f0fdf4", "#166534"),
            "Packed": ("#f0fdf4", "#166534"),
            "Ready for Dispatch": ("#fefce8", "#854d0e"),
            "Dispatched": ("#fefce8", "#854d0e"),
            "Delivered": ("#ecfdf5", "#065f46"),
            "Returned": ("#fff1f2", "#9f1239"),
            "Cancelled": ("#f3f4f6", "#6b7280"),
        }
        for status, (bg, fg) in palette.items():
            tag = f"status::{self._norm(status)}"
            self.outbound_summary_tree.tag_configure(tag, background=bg, foreground=fg)
        self.outbound_summary_tree.tag_configure("zebra::odd", background="#ffffff")
        self.outbound_summary_tree.tag_configure("zebra::even", background="#f8fafc")
        self._outbound_tags_ready = True

    def _outbound_status_tag(self, status):
        return f"status::{self._norm(status)}" if self._clean_text(status) else ""

    def _populate_outbound_summary_tree(self):
        self._ensure_outbound_status_tags()
        for child in self.outbound_summary_tree.get_children():
            self.outbound_summary_tree.delete(child)

        mode = self._outbound_view_mode()
        self._configure_outbound_summary_columns(mode)
        total_rows = len(self._outbound_view_rows)
        try:
            page_size = int(self.outbound_page_size)
        except Exception:
            page_size = 25
        page_size = max(1, page_size)
        max_page = max(1, int((total_rows + page_size - 1) / page_size))
        self.outbound_page = max(1, min(self.outbound_page, max_page))
        start = (self.outbound_page - 1) * page_size
        page_rows = self._outbound_view_rows[start:start + page_size]
        self.outbound_page_label.set(f"Page {self.outbound_page} / {max_page} (Rows {start + 1}-{min(start + page_size, total_rows)} of {total_rows})" if total_rows else "Page 1 / 1")

        self._outbound_expand_state = {}
        for idx, row in enumerate(page_rows):
            status = self._clean_text(row.get("status"))
            iid = f"view::{idx}::{self._clean_text(row.get('key'))}"
            self._outbound_view_map[iid] = row
            zebra_tag = "zebra::even" if idx % 2 == 0 else "zebra::odd"
            if mode == "Per Item":
                self.outbound_summary_tree.insert(
                    "",
                    "end",
                    iid=iid,
                    text="",
                    values=(
                        row.get("outbound_no"),
                        row.get("date"),
                        row.get("client"),
                        row.get("destination"),
                        row.get("warehouse"),
                        row.get("item_code"),
                        row.get("product_name"),
                        row.get("batch_number"),
                        self._format_quantity(row.get("qty") or 0),
                        status,
                    ),
                    tags=(zebra_tag, self._outbound_status_tag(status)),
                )
            elif mode == "Per Destination":
                self.outbound_summary_tree.insert(
                    "",
                    "end",
                    iid=iid,
                    text="",
                    values=(
                        row.get("destination"),
                        row.get("client"),
                        row.get("warehouse"),
                        row.get("shipments"),
                        row.get("num_products"),
                        self._format_quantity(row.get("total_qty") or 0),
                        status,
                    ),
                    tags=(zebra_tag, self._outbound_status_tag(status)),
                )
            else:
                self.outbound_summary_tree.insert(
                    "",
                    "end",
                    iid=iid,
                    text="▸",
                    values=(
                        row.get("outbound_no"),
                        row.get("date"),
                        row.get("client"),
                        row.get("destination"),
                        row.get("warehouse"),
                        row.get("num_products"),
                        self._format_quantity(row.get("total_qty") or 0),
                        status,
                    ),
                    tags=(zebra_tag, self._outbound_status_tag(status)),
                )

    def _toggle_outbound_expand(self, _event=None):
        if self._outbound_view_mode() != "Per Shipment":
            return
        sel = self.outbound_summary_tree.selection()
        if not sel:
            return
        item_id = sel[0]
        view = self._outbound_view_map.get(item_id)
        if not view:
            return
        outbound_no = self._clean_text(view.get("outbound_no"))
        if not outbound_no:
            return
        expanded = bool(self._outbound_expand_state.get(item_id))
        if expanded:
            for child in self.outbound_summary_tree.get_children(item_id):
                self.outbound_summary_tree.delete(child)
            self.outbound_summary_tree.item(item_id, text="▸")
            self._outbound_expand_state[item_id] = False
            return

        for idx, it in enumerate(view.get("_items") or []):
            self.outbound_summary_tree.insert(
                item_id,
                "end",
                iid=f"{item_id}::line::{idx}",
                text="",
                values=(
                    it.get("item_code"),
                    it.get("product_name"),
                    it.get("brand"),
                    it.get("sku"),
                    it.get("batch_number"),
                    "Item Detail",
                    it.get("quantity_dispatched"),
                    it.get("expiry_date"),
                ),
            )
        self.outbound_summary_tree.item(item_id, open=True, text="▾")
        self._outbound_expand_state[item_id] = True

    def _on_outbound_row_selected(self, _event=None):
        sel = self.outbound_summary_tree.selection()
        if not sel:
            return
        node = sel[0]
        view = self._outbound_view_map.get(node)
        if not view:
            parent = self.outbound_summary_tree.parent(node)
            view = self._outbound_view_map.get(parent)
        if not view:
            return
        self._populate_outbound_details_from_view(view)

    def _populate_outbound_details_from_view(self, view):
        for child in self.outbound_details_tree.get_children():
            self.outbound_details_tree.delete(child)
        if not view:
            return
        items = view.get("_items") or []
        shipments = view.get("_shipments") or []
        for it in items:
            self.outbound_details_tree.insert(
                "",
                "end",
                values=(
                    it.get("item_code"),
                    it.get("product_name"),
                    it.get("brand"),
                    it.get("sku"),
                    it.get("batch_number"),
                    it.get("lot_number"),
                    it.get("manufacture_date"),
                    it.get("expiry_date"),
                    it.get("unit_pack_size"),
                    it.get("quantity_dispatched"),
                    it.get("warehouse"),
                    it.get("bin_location"),
                    it.get("temperature_zone"),
                    it.get("notes"),
                ),
            )
        if len(shipments) == 1:
            row = shipments[0]
            self.outbound_dispatch_vars["requested_by"].set(row.get("requested_by") or "-")
            self.outbound_dispatch_vars["approved_by"].set(row.get("approved_by") or "-")
            self.outbound_dispatch_vars["picked_by"].set(row.get("picked_by") or "-")
            self.outbound_dispatch_vars["packed_by"].set(row.get("packed_by") or "-")
            self.outbound_dispatch_vars["dispatched_by"].set(row.get("dispatched_by") or "-")
            self.outbound_dispatch_vars["received_by"].set(row.get("received_by") or "-")
            self.outbound_dispatch_vars["dispatch_date"].set(row.get("dispatch_date") or row.get("date") or "-")
            self.outbound_dispatch_vars["delivery_date"].set(row.get("delivery_date") or "-")
            self.outbound_dispatch_vars["delivery_note_no"].set(row.get("delivery_note_no") or "-")
            self.outbound_dispatch_vars["goods_transfer_note_no"].set(row.get("delivery_note_no") or "-")
            self.outbound_dispatch_vars["vehicle"].set(row.get("vehicle") or "-")
            self.outbound_dispatch_vars["courier"].set(row.get("courier") or "-")
            self.outbound_dispatch_vars["tracking_number"].set(row.get("tracking_number") or "-")
            self._selected_outbound_no = self._clean_text(row.get("outbound_no"))
        else:
            self.outbound_dispatch_vars["requested_by"].set("Multiple")
            self.outbound_dispatch_vars["approved_by"].set("Multiple")
            self.outbound_dispatch_vars["picked_by"].set("Multiple")
            self.outbound_dispatch_vars["packed_by"].set("Multiple")
            self.outbound_dispatch_vars["dispatched_by"].set("Multiple")
            self.outbound_dispatch_vars["received_by"].set("Multiple")
            self.outbound_dispatch_vars["dispatch_date"].set("Multiple")
            self.outbound_dispatch_vars["delivery_date"].set("Multiple")
            self.outbound_dispatch_vars["delivery_note_no"].set("Multiple")
            self.outbound_dispatch_vars["goods_transfer_note_no"].set("Multiple")
            self.outbound_dispatch_vars["vehicle"].set("-")
            self.outbound_dispatch_vars["courier"].set("-")
            self.outbound_dispatch_vars["tracking_number"].set("-")
            self._selected_outbound_no = ""

    def _configure_outbound_summary_columns(self, mode):
        if getattr(self, "_outbound_columns_mode", None) == mode:
            return
        if mode == "Per Item":
            columns = ("Outbound Number", "Date", "Client", "Destination", "Warehouse", "Item Code", "Product", "Batch", "Quantity", "Status")
        elif mode == "Per Destination":
            columns = ("Destination", "Client", "Warehouse", "Shipments", "Number of Products", "Total Quantity", "Status")
        else:
            columns = ("Outbound Number", "Date", "Client", "Destination", "Warehouse", "Number of Products", "Total Quantity", "Status")
        self.outbound_summary_tree["columns"] = columns
        for col in columns:
            self.outbound_summary_tree.heading(col, text=col, command=lambda c=col: self._sort_outbound_summary(c))
            self.outbound_summary_tree.column(col, width=140, stretch=True)
        self._outbound_columns_mode = mode

    def _populate_outbound_traceability(self):
        for child in self.outbound_trace_tree.get_children():
            self.outbound_trace_tree.delete(child)
        for row in self.outbound_filtered:
            for it in row.get("items") or []:
                self.outbound_trace_tree.insert(
                    "",
                    "end",
                    values=(
                        row.get("client") or "-",
                        row.get("warehouse") or "-",
                        it.get("product_name") or "-",
                        it.get("batch_number") or "-",
                        it.get("lot_number") or "-",
                        it.get("manufacture_date") or "-",
                        it.get("expiry_date") or "-",
                        it.get("quantity_dispatched") or "-",
                        row.get("dispatch_date") or row.get("date") or "-",
                        row.get("destination") or "-",
                    ),
                )

    def _prev_outbound_page(self):
        if self.outbound_page > 1:
            self.outbound_page -= 1
            self._populate_outbound_summary_tree()

    def _next_outbound_page(self):
        total_rows = len(self.outbound_filtered)
        try:
            page_size = int(self.outbound_page_size)
        except Exception:
            page_size = 25
        max_page = max(1, int((total_rows + page_size - 1) / page_size))
        if self.outbound_page < max_page:
            self.outbound_page += 1
            self._populate_outbound_summary_tree()

    def _apply_outbound_page_size(self):
        try:
            self.outbound_page_size = max(1, int(self.outbound_page_size_var.get() or 25))
        except Exception:
            self.outbound_page_size = 25
        self.outbound_page = 1
        self._populate_outbound_summary_tree()

    def _sort_outbound_summary(self, column):
        current = getattr(self, "_outbound_sort", ("Date", True))
        asc = True
        if current[0] == column:
            asc = not bool(current[1])
        self._outbound_sort = (column, asc)

        mode = self._outbound_view_mode()

        def key_fn(row):
            if mode == "Per Item":
                if column == "Outbound Number":
                    return self._clean_text(row.get("outbound_no"))
                if column == "Date":
                    return self._sort_key(row.get("date"))
                if column == "Client":
                    return self._clean_text(row.get("client")).lower()
                if column == "Destination":
                    return self._clean_text(row.get("destination")).lower()
                if column == "Warehouse":
                    return self._clean_text(row.get("warehouse")).lower()
                if column == "Item Code":
                    return self._clean_text(row.get("item_code")).lower()
                if column == "Product":
                    return self._clean_text(row.get("product_name")).lower()
                if column == "Batch":
                    return self._clean_text(row.get("batch_number")).lower()
                if column == "Quantity":
                    return float(row.get("qty") or 0)
                if column == "Status":
                    return self._clean_text(row.get("status")).lower()
                return ""

            if mode == "Per Destination":
                if column == "Destination":
                    return self._clean_text(row.get("destination")).lower()
                if column == "Client":
                    return self._clean_text(row.get("client")).lower()
                if column == "Warehouse":
                    return self._clean_text(row.get("warehouse")).lower()
                if column == "Shipments":
                    return int(row.get("shipments") or 0)
                if column == "Number of Products":
                    return int(row.get("num_products") or 0)
                if column == "Total Quantity":
                    return float(row.get("total_qty") or 0)
                if column == "Status":
                    return self._clean_text(row.get("status")).lower()
                return ""

            if column == "Outbound Number":
                return self._clean_text(row.get("outbound_no"))
            if column == "Date":
                return self._sort_key(row.get("date"))
            if column == "Client":
                return self._clean_text(row.get("client")).lower()
            if column == "Destination":
                return self._clean_text(row.get("destination")).lower()
            if column == "Warehouse":
                return self._clean_text(row.get("warehouse")).lower()
            if column == "Number of Products":
                return int(row.get("num_products") or 0)
            if column == "Total Quantity":
                return float(row.get("total_qty") or 0)
            if column == "Status":
                return self._clean_text(row.get("status")).lower()
            return ""

        self._outbound_view_rows.sort(key=key_fn, reverse=not asc)
        self.outbound_page = 1
        self._populate_outbound_summary_tree()

    def _selected_outbound_for_docs(self):
        outbound_no = self._clean_text(getattr(self, "_selected_outbound_no", ""))
        if outbound_no:
            return outbound_no
        sel = self.outbound_summary_tree.selection()
        if not sel:
            return ""
        node = sel[0]
        view = getattr(self, "_outbound_view_map", {}).get(node)
        if not view:
            view = getattr(self, "_outbound_view_map", {}).get(self.outbound_summary_tree.parent(node))
        if not view:
            return ""
        shipments = view.get("_shipments") or []
        if len(shipments) != 1:
            return ""
        return self._clean_text(shipments[0].get("outbound_no"))

    def _open_outbound_pick_list(self):
        outbound_no = self._selected_outbound_for_docs()
        if not outbound_no:
            messagebox.showwarning("Pick List", "Select an outbound record first.")
            return
        try:
            mgr = WarehouseManager(self.data_folder)
            mgr.export_pick_list_pdf(outbound_no, open_after=True)
        except Exception as exc:
            messagebox.showerror("Pick List", str(exc))

    def _open_outbound_packing_list(self):
        outbound_no = self._selected_outbound_for_docs()
        if not outbound_no:
            messagebox.showwarning("Packing List", "Select an outbound record first.")
            return
        try:
            mgr = WarehouseManager(self.data_folder)
            mgr.export_packing_list_pdf(outbound_no, open_after=True)
        except Exception as exc:
            messagebox.showerror("Packing List", str(exc))

    def _open_outbound_delivery_note(self):
        outbound_no = self._selected_outbound_for_docs()
        if not outbound_no:
            messagebox.showwarning("Delivery Note", "Select an outbound record first.")
            return
        try:
            mgr = WarehouseManager(self.data_folder)
            mgr.export_outbound_delivery_note_pdf(outbound_no, open_after=True)
        except Exception as exc:
            messagebox.showerror("Delivery Note", str(exc))

    def _open_outbound_pod(self):
        outbound_no = self._selected_outbound_for_docs()
        if not outbound_no:
            messagebox.showwarning("Proof of Delivery", "Select an outbound record first.")
            return
        row = next((r for r in self.outbound_filtered if self._clean_text(r.get("outbound_no")) == self._clean_text(outbound_no)), None)
        path = self._clean_text((row or {}).get("proof_of_delivery_path"))
        if path and os.path.exists(path):
            try:
                webbrowser.open("file://" + os.path.abspath(path))
            except Exception:
                pass
            return
        messagebox.showinfo("Proof of Delivery", "No Proof of Delivery file is attached for this outbound record.")

    def _outbound_export_rows(self):
        rows = []
        for outbound in self.outbound_filtered:
            for it in outbound.get("items") or []:
                rows.append({
                    "Outbound Number": outbound.get("outbound_no"),
                    "Delivery Note Number": outbound.get("delivery_note_no"),
                    "Date": outbound.get("date"),
                    "Client": outbound.get("client"),
                    "Destination": outbound.get("destination"),
                    "Warehouse": outbound.get("warehouse"),
                    "Bin Location": it.get("bin_location") or outbound.get("bin_location"),
                    "Item Code": it.get("item_code"),
                    "Product Name": it.get("product_name"),
                    "Brand": it.get("brand"),
                    "SKU": it.get("sku"),
                    "Batch Number": it.get("batch_number"),
                    "Lot Number": it.get("lot_number"),
                    "Manufacturing Date": it.get("manufacture_date"),
                    "Expiry Date": it.get("expiry_date"),
                    "Unit / Pack Size": it.get("unit_pack_size"),
                    "Quantity Dispatched": it.get("quantity_dispatched"),
                    "Temperature Zone": it.get("temperature_zone"),
                    "Status": outbound.get("status"),
                    "Requested By": outbound.get("requested_by"),
                    "Approved By": outbound.get("approved_by"),
                    "Picked By": outbound.get("picked_by"),
                    "Packed By": outbound.get("packed_by"),
                    "Dispatched By": outbound.get("dispatched_by"),
                    "Received By": outbound.get("received_by"),
                    "Dispatch Date": outbound.get("dispatch_date"),
                    "Delivery Date": outbound.get("delivery_date"),
                    "Notes": outbound.get("notes"),
                })
        return rows

    def export_outbound_report_csv(self):
        rows = self._outbound_export_rows()
        if not rows:
            messagebox.showwarning("Export CSV", "No outbound rows to export.")
            return
        folder = os.path.join(self.data_folder, "InventoryReports")
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"Warehouse_Outbound_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv")
        try:
            with open(path, "w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
                writer.writeheader()
                writer.writerows(rows)
            webbrowser.open("file://" + os.path.abspath(path))
        except Exception as exc:
            messagebox.showerror("Export CSV", f"Failed to export CSV:\n{exc}")
            return
        messagebox.showinfo("Export CSV", f"CSV saved:\n{path}")

    def export_outbound_report_excel(self):
        rows = self._outbound_export_rows()
        if not rows:
            messagebox.showwarning("Export Excel", "No outbound rows to export.")
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except Exception:
            messagebox.showerror("Export Excel", "Excel export is not available in this build.")
            return

        folder = os.path.join(self.data_folder, "InventoryReports")
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"Warehouse_Outbound_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Outbound Report"
            headers = list(rows[0].keys())
            ws.append(headers)
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1F2937")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in rows:
                ws.append([row.get(h, "") for h in headers])
            ws.freeze_panes = "A2"
            for col_idx, header in enumerate(headers, start=1):
                max_len = len(str(header))
                for row_idx in range(2, min(ws.max_row, 800) + 1):
                    v = ws.cell(row=row_idx, column=col_idx).value
                    if v is None:
                        continue
                    max_len = max(max_len, len(str(v)))
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max(12, max_len + 2), 55)
            wb.save(path)
            webbrowser.open("file://" + os.path.abspath(path))
        except Exception as exc:
            messagebox.showerror("Export Excel", f"Failed to export Excel:\n{exc}")
            return
        messagebox.showinfo("Export Excel", f"Excel saved:\n{path}")

    def export_outbound_report_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Export PDF", "PDF export is not available in this build.")
            return
        rows = list(self.outbound_filtered or [])
        if not rows:
            messagebox.showwarning("Export PDF", "No outbound rows to export.")
            return
        folder = os.path.join(self.data_folder, "InventoryReports")
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"Warehouse_Outbound_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        try:
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle("OutboundTitle", parent=styles["Heading1"], fontSize=16, alignment=1, textColor=colors.HexColor("#0b5394"))
            sub_style = ParagraphStyle("OutboundSub", parent=styles["Normal"], fontSize=9, alignment=1, textColor=colors.HexColor("#64748b"))
            info_style = ParagraphStyle("OutboundInfo", parent=styles["Normal"], fontSize=8.5, alignment=1, textColor=colors.HexColor("#475569"))
            cell_style = ParagraphStyle("OutboundCell", parent=styles["BodyText"], fontSize=7.8, leading=9, wordWrap="CJK")

            doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=0.35 * inch, rightMargin=0.35 * inch, topMargin=0.4 * inch, bottomMargin=0.45 * inch)
            elements = []
            elements.extend(build_pdf_logo_flowables(self.data_folder, width=1.0 * inch, height=1.0 * inch, spacer_height=0.08 * inch))

            filters = self._outbound_current_filters()
            elements.append(Paragraph("Warehouse Outbound Report", title_style))
            elements.append(Paragraph("Warehouse dispatch & batch traceability report (no financial information).", sub_style))
            elements.append(Paragraph(
                f"Date From: {filters.get('date_from') or '-'} | Date To: {filters.get('date_to') or '-'} | Warehouse: {filters.get('warehouse') or 'All'} | Client: {filters.get('client') or 'All'} | Status: {filters.get('status') or 'All'}",
                info_style,
            ))
            elements.append(Paragraph(f"Generated By: {getpass_user_fallback()} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", info_style))
            elements.append(Spacer(1, 10))

            summary_pairs = [
                ("Total Outbound Orders", self.outbound_summary_vars["total_orders"].get()),
                ("Total Clients", self.outbound_summary_vars["total_clients"].get()),
                ("Total Products Dispatched", self.outbound_summary_vars["total_products"].get()),
                ("Total Quantity Dispatched", self.outbound_summary_vars["total_qty"].get()),
                ("Completed Orders", self.outbound_summary_vars["completed"].get()),
                ("Pending Orders", self.outbound_summary_vars["pending"].get()),
                ("Returned Orders", self.outbound_summary_vars["returned"].get()),
                ("Cancelled Orders", self.outbound_summary_vars["cancelled"].get()),
            ]
            summary_rows = []
            for i in range(0, len(summary_pairs), 4):
                chunk = summary_pairs[i:i + 4]
                row_cells = []
                for label, val in chunk:
                    row_cells.append(Paragraph(f"<b>{label}</b><br/>{self._clean_text(val)}", cell_style))
                summary_rows.append(row_cells)
            summary_table = Table(summary_rows, colWidths=[doc.width / 4.0] * 4)
            summary_table.setStyle(TableStyle([
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e2e8f0")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            elements.append(summary_table)
            elements.append(Spacer(1, 10))

            status_priority = [
                "Returned",
                "Cancelled",
                "Delivered",
                "Dispatched",
                "Ready for Dispatch",
                "Packed",
                "Packing",
                "Picked",
                "Picking",
                "Reserved",
                "Approved",
                "Pending Approval",
                "Draft",
            ]
            priority_map = {self._norm(x): i for i, x in enumerate(status_priority)}

            grouped = {}
            for ob in rows:
                key = (self._norm(ob.get("client")), self._norm(ob.get("destination")))
                grouped.setdefault(key, []).append(ob)

            grouped_rows = []
            for (_c, _d), group in grouped.items():
                client_name = self._clean_text(group[0].get("client")) if group else ""
                destination_name = self._clean_text(group[0].get("destination")) if group else ""
                warehouses = set()
                shipment_ids = []
                product_count = 0
                total_qty = 0.0
                best_status = None
                best_status_label = ""
                latest_date = ""
                for ob in group:
                    shipment_ids.append(self._clean_text(ob.get("outbound_no")))
                    warehouses.add(self._clean_text(ob.get("warehouse")))
                    ob_status = self._clean_text(ob.get("status"))
                    idx = priority_map.get(self._norm(ob_status))
                    if idx is not None and (best_status is None or idx < best_status):
                        best_status = idx
                        best_status_label = ob_status
                    ob_date = self._clean_text(ob.get("date"))
                    if ob_date:
                        if not latest_date or self._sort_key(ob_date) > self._sort_key(latest_date):
                            latest_date = ob_date
                    items = ob.get("items") or []
                    product_count += len(items)
                    for it in items:
                        try:
                            total_qty += float(it.get("raw_quantity_dispatched") or 0)
                        except Exception:
                            pass
                grouped_rows.append({
                    "client": client_name,
                    "destination": destination_name,
                    "date": latest_date,
                    "shipments": sorted([x for x in shipment_ids if x]),
                    "warehouse_list": sorted([x for x in warehouses if x]),
                    "num_products": product_count,
                    "total_qty": total_qty,
                    "status": best_status_label or "Draft",
                    "records": group,
                })

            grouped_rows.sort(key=lambda r: self._sort_key(r.get("date")), reverse=True)

            summary_headers = ["Client", "Destination", "Last Date", "Shipments", "Warehouses", "No. Products", "Total Qty", "Status"]
            summary_table_rows = [summary_headers]
            for grp in grouped_rows:
                summary_table_rows.append([
                    self._pdf_paragraph(grp.get("client"), cell_style),
                    self._pdf_paragraph(grp.get("destination"), cell_style),
                    self._pdf_paragraph(grp.get("date"), cell_style),
                    self._pdf_paragraph(str(len(grp.get("shipments") or [])), cell_style),
                    self._pdf_paragraph(", ".join(grp.get("warehouse_list") or []) or "-", cell_style),
                    self._pdf_paragraph(str(grp.get("num_products") or 0), cell_style),
                    self._pdf_paragraph(self._format_quantity(grp.get("total_qty") or 0), cell_style),
                    self._pdf_paragraph(grp.get("status"), cell_style),
                ])
            summ_table = Table(summary_table_rows, colWidths=[95, 120, 70, 60, 90, 60, 70, 70], repeatRows=1)
            summ_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b5394")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            elements.append(Paragraph("Outbound Summary (Per Shipment Grouped by Client + Destination)", ParagraphStyle("Sec", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#111827"))))
            elements.append(summ_table)
            elements.append(Spacer(1, 10))

            elements.append(Paragraph("Detailed Shipments (Per Group)", ParagraphStyle("Sec2", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#111827"))))
            for grp in grouped_rows:
                elements.append(Spacer(1, 6))
                elements.append(Paragraph(
                    f"<b>Client:</b> {self._clean_text(grp.get('client')) or '-'} &nbsp;&nbsp; <b>Destination:</b> {self._clean_text(grp.get('destination')) or '-'}",
                    info_style,
                ))
                shipment_headers = ["Shipment No.", "Date", "Warehouse", "Status", "Delivery Note"]
                shipment_rows = [shipment_headers]
                for ob in grp.get("records") or []:
                    shipment_rows.append([
                        self._pdf_paragraph(ob.get("outbound_no"), cell_style),
                        self._pdf_paragraph(ob.get("date"), cell_style),
                        self._pdf_paragraph(ob.get("warehouse"), cell_style),
                        self._pdf_paragraph(ob.get("status"), cell_style),
                        self._pdf_paragraph(ob.get("delivery_note_no"), cell_style),
                    ])
                ship_table = Table(shipment_rows, colWidths=[90, 70, 70, 70, 85], repeatRows=1)
                ship_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17395d")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                elements.append(ship_table)

                trace_headers = ["Product", "Batch", "Lot", "MFG", "EXP", "Qty", "Warehouse", "Bin"]
                trace_rows = [trace_headers]
                for ob in grp.get("records") or []:
                    for it in ob.get("items") or []:
                        trace_rows.append([
                            self._pdf_paragraph(it.get("product_name"), cell_style),
                            self._pdf_paragraph(it.get("batch_number"), cell_style),
                            self._pdf_paragraph(it.get("lot_number"), cell_style),
                            self._pdf_paragraph(it.get("manufacture_date"), cell_style),
                            self._pdf_paragraph(it.get("expiry_date"), cell_style),
                            self._pdf_paragraph(it.get("quantity_dispatched"), cell_style),
                            self._pdf_paragraph(it.get("warehouse") or ob.get("warehouse"), cell_style),
                            self._pdf_paragraph(it.get("bin_location") or ob.get("bin_location"), cell_style),
                        ])
                detail_table = Table(trace_rows, colWidths=[140, 70, 60, 55, 55, 45, 70, 70], repeatRows=1)
                detail_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b5394")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                elements.append(Spacer(1, 6))
                elements.append(detail_table)
                elements.append(Spacer(1, 8))

            footer_style = ParagraphStyle("Footer", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#475569"))
            elements.append(Paragraph("Prepared By: ____________________    Reviewed By: ____________________    Approved By: ____________________", footer_style))
            elements.append(Paragraph(f"Report Generation Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", footer_style))

            def _page(canvas, _doc):
                canvas.saveState()
                canvas.setFont("Helvetica", 8)
                canvas.setFillColor(colors.HexColor("#64748b"))
                canvas.drawRightString(_doc.pagesize[0] - 0.4 * inch, 0.25 * inch, f"Page {_doc.page}")
                canvas.restoreState()

            doc.build(elements, onFirstPage=_page, onLaterPages=_page)
        except Exception as exc:
            messagebox.showerror("Export PDF", f"Failed to generate PDF:\n{exc}")
            return
        self._open_exported_pdf(path)
