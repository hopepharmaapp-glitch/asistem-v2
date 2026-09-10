import json
import os
import csv
import getpass
import re
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext

TEMP_UNIT = " °C"


def _strip_temperature_unit(value):
    text = str(value or "").strip()
    if not text:
        return ""
    lowered = text.lower().replace(" ", "")
    if lowered.endswith("°c"):
        return text[:-2].replace("°", "").strip()
    return text


def _truncate_temperature_string(value):
    text = _strip_temperature_unit(value)
    if not text:
        return ""
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
        raise ValueError("Temperature must be a number")
    sign = ""
    body = text
    if body[0] in "+-":
        sign = body[0]
        body = body[1:]
    if "." not in body:
        return f"{sign}{body}"
    whole, fraction = body.split(".", 1)
    if not fraction:
        return f"{sign}{whole}"
    return f"{sign}{whole}.{fraction[:1]}"


def format_temperature_display(value):
    try:
        truncated = _truncate_temperature_string(value)
    except Exception:
        return str(value or "").strip()
    if not truncated:
        return ""
    try:
        numeric = float(truncated)
    except Exception:
        return str(value or "").strip()
    if numeric.is_integer():
        return f"{int(numeric)}{TEMP_UNIT}"
    return f"{truncated}{TEMP_UNIT}"


def parse_temperature_value(value, field_name="Temperature"):
    try:
        truncated = _truncate_temperature_string(value)
    except Exception:
        raise ValueError(f"{field_name} must be a valid temperature")
    if not truncated:
        raise ValueError(f"{field_name} is required")
    try:
        return float(truncated)
    except Exception:
        raise ValueError(f"{field_name} must be a valid temperature")


class TemperatureLogManager:
    DEFAULT_POINTS = [f"Point {i}" for i in range(1, 7)]
    DEFAULT_SLOTS = ["Morning", "Evening"]

    def __init__(self, data_folder):
        self.data_folder = data_folder
        self.logs_file = os.path.join(data_folder, "temperature_logs.json")
        self.audit_log_file = os.path.join(data_folder, "audit_log.jsonl")
        self.settings_file = os.path.join(data_folder, "app_settings.json")
        self._ensure_temperature_settings()
        self.logs = self.load_logs()

    def load_logs(self):
        try:
            if os.path.exists(self.logs_file):
                with open(self.logs_file, "r", encoding="utf-8") as f:
                    data = json.load(f) or []
                if isinstance(data, list):
                    rows = []
                    for row in data:
                        if not isinstance(row, dict):
                            continue
                        row.setdefault("point_name", "")
                        row.setdefault("slot_name", "")
                        row.setdefault("notes", "")
                        row.setdefault("recorded_by", "")
                        rows.append(row)
                    return rows
        except Exception:
            pass
        return []

    def save_logs(self):
        try:
            os.makedirs(self.data_folder, exist_ok=True)
            with open(self.logs_file, "w", encoding="utf-8") as f:
                json.dump(list(self.logs or []), f, indent=2, ensure_ascii=False)
            return True
        except Exception:
            return False

    def _load_settings(self):
        try:
            if os.path.exists(self.settings_file):
                with open(self.settings_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                return data if isinstance(data, dict) else {}
        except Exception:
            pass
        return {}

    def _save_settings(self, settings_dict):
        try:
            if not isinstance(settings_dict, dict):
                return False
            os.makedirs(self.data_folder, exist_ok=True)
            with open(self.settings_file, "w", encoding="utf-8") as f:
                json.dump(settings_dict, f, indent=2, ensure_ascii=False)
            return True
        except Exception:
            return False

    def _ensure_temperature_settings(self):
        settings = self._load_settings()
        changed = False

        pts = settings.get("temperature_points")
        if not isinstance(pts, list):
            pts = []
        cleaned_pts = [str(x or "").strip() for x in pts if str(x or "").strip()]
        if len(cleaned_pts) < 6:
            cleaned_pts.extend(self.DEFAULT_POINTS[len(cleaned_pts):6])
            changed = True
        elif len(cleaned_pts) > 6:
            cleaned_pts = cleaned_pts[:6]
            changed = True
        if settings.get("temperature_points") != cleaned_pts:
            settings["temperature_points"] = cleaned_pts
            changed = True

        if settings.get("temperature_slots") != list(self.DEFAULT_SLOTS):
            settings["temperature_slots"] = list(self.DEFAULT_SLOTS)
            changed = True

        if changed:
            self._save_settings(settings)

    def get_temperature_points(self):
        settings = self._load_settings()
        pts = settings.get("temperature_points")
        if isinstance(pts, list):
            cleaned = [str(x or "").strip() for x in pts if str(x or "").strip()]
            if len(cleaned) >= 6:
                return cleaned[:6]
        return list(self.DEFAULT_POINTS)

    def set_temperature_points(self, point_names):
        if not isinstance(point_names, list):
            return False, "Invalid point list"
        cleaned = [str(x or "").strip() for x in point_names]
        if len(cleaned) != 6 or any(not x for x in cleaned):
            return False, "Exactly 6 point names are required"
        settings = self._load_settings()
        settings["temperature_points"] = cleaned
        settings["temperature_slots"] = list(self.DEFAULT_SLOTS)
        if not self._save_settings(settings):
            return False, "Failed to save settings"
        return True, "Saved"

    def get_temperature_slots(self):
        return list(self.DEFAULT_SLOTS)

    def append_audit(self, event, details):
        try:
            entry = {
                "ts": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "event": str(event or "").strip(),
                "user": getpass.getuser(),
                "details": details if isinstance(details, dict) else {"value": str(details)},
            }
            with open(self.audit_log_file, "a", encoding="utf-8") as f:
                f.write(json.dumps(entry, ensure_ascii=False) + "\n")
            return True
        except Exception:
            return False

    def _new_id(self):
        return f"TLOG{datetime.now().strftime('%Y%m%d%H%M%S%f')}"

    def _require_non_empty(self, value, field_name):
        s = str(value or "").strip()
        if not s:
            raise ValueError(f"{field_name} is required")
        return s

    def _require_float(self, value, field_name):
        try:
            return float(value)
        except Exception:
            raise ValueError(f"{field_name} must be a number")

    def _require_temperature(self, value, field_name):
        return parse_temperature_value(value, field_name)

    def _parse_datetime(self, dt_value):
        s = str(dt_value or "").strip()
        if not s:
            return datetime.now()
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(s, fmt)
            except Exception:
                pass
        try:
            return datetime.fromisoformat(s)
        except Exception:
            return datetime.now()

    def get_date_key(self, timestamp):
        return self._parse_datetime(timestamp).strftime("%Y-%m-%d")

    def _make_session_key(self, timestamp, room_name, slot_name):
        ts = self._parse_datetime(timestamp).strftime("%Y-%m-%d %H:%M:%S")
        return "||".join([ts, str(room_name or "").strip(), str(slot_name or "").strip()])

    def _session_key_from_row(self, row):
        return self._make_session_key(row.get("timestamp"), row.get("room_name"), row.get("slot_name"))

    def session_exists(self, date_key, room_name, slot_name, exclude_session_key=None):
        date_key = str(date_key or "").strip()
        room_name = str(room_name or "").strip()
        slot_name = str(slot_name or "").strip()
        exclude_session_key = str(exclude_session_key or "").strip() or None
        for row in (self.logs or []):
            try:
                if self.get_date_key(row.get("timestamp")) != date_key:
                    continue
                if str(row.get("room_name") or "").strip() != room_name:
                    continue
                if str(row.get("slot_name") or "").strip() != slot_name:
                    continue
                session_key = self._session_key_from_row(row)
                if exclude_session_key and session_key == exclude_session_key:
                    continue
                return True
            except Exception:
                pass
        return False

    def _normalize_point_readings(self, point_readings):
        if not isinstance(point_readings, dict):
            raise ValueError("Point readings are required")
        out = {}
        for point_name in self.get_temperature_points():
            raw = str(point_readings.get(point_name, "") or "").strip()
            if not raw:
                raise ValueError(f"{point_name} temperature is required")
            out[point_name] = self._require_temperature(raw, f"{point_name} temperature")
        return out

    def add_session(self, timestamp, room_name, slot_name, point_readings, notes="", recorded_by=""):
        ts = self._parse_datetime(timestamp)
        room = self._require_non_empty(room_name, "Room name")
        slot = self._require_non_empty(slot_name, "Slot")
        if slot not in self.get_temperature_slots():
            raise ValueError("Slot must be Morning or Evening")
        readings = self._normalize_point_readings(point_readings)
        user_name = self._require_non_empty(recorded_by or getpass.getuser(), "User")

        entries = []
        for point_name, value in readings.items():
            entries.append({
                "log_id": self._new_id(),
                "timestamp": ts.strftime("%Y-%m-%d %H:%M:%S"),
                "room_name": room,
                "point_name": point_name,
                "slot_name": slot,
                "current_temp": value,
                "min_temp": value,
                "max_temp": value,
                "notes": str(notes or "").strip(),
                "recorded_by": user_name,
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })

        previous = list(self.logs or [])
        self.logs = previous + entries
        if not self.save_logs():
            self.logs = previous
            return False, "Failed to save temperature log", None

        session_key = self._make_session_key(ts, room, slot)
        self.append_audit(
            "temperature_log_session_created",
            {"session_key": session_key, "room": room, "slot": slot, "points": len(entries), "recorded_by": user_name}
        )
        return True, "Temperature log saved", session_key

    def update_session(self, session_key, timestamp, room_name, slot_name, point_readings, notes="", recorded_by=""):
        old_entries = [row for row in (self.logs or []) if self._session_key_from_row(row) == session_key]
        if not old_entries:
            return False, "Temperature log session not found", None
        keep = [row for row in (self.logs or []) if self._session_key_from_row(row) != session_key]
        snapshot = list(self.logs or [])
        self.logs = keep
        ok, msg, new_session_key = self.add_session(timestamp, room_name, slot_name, point_readings, notes=notes, recorded_by=recorded_by)
        if not ok:
            self.logs = snapshot
            self.save_logs()
            return False, msg, None
        self.append_audit(
            "temperature_log_session_updated",
            {"old_session_key": session_key, "new_session_key": new_session_key}
        )
        return True, "Temperature log updated", new_session_key

    def delete_session(self, session_key):
        before = len(self.logs or [])
        self.logs = [row for row in (self.logs or []) if self._session_key_from_row(row) != session_key]
        if len(self.logs or []) == before:
            return False, "Temperature log session not found"
        if not self.save_logs():
            return False, "Failed to save deletion"
        self.append_audit("temperature_log_session_deleted", {"session_key": session_key})
        return True, "Temperature log deleted"

    def get_grouped_logs(self, start_date=None, end_date=None, room_query=""):
        query = str(room_query or "").strip().lower()
        start_dt = None
        end_dt = None
        if start_date:
            try:
                start_dt = datetime.strptime(str(start_date).strip(), "%Y-%m-%d")
            except Exception:
                pass
        if end_date:
            try:
                end_dt = datetime.strptime(str(end_date).strip(), "%Y-%m-%d")
            except Exception:
                pass

        grouped = {}
        for row in (self.logs or []):
            try:
                ts = self._parse_datetime(row.get("timestamp"))
                if start_dt and ts < start_dt:
                    continue
                if end_dt and ts > end_dt.replace(hour=23, minute=59, second=59):
                    continue
                if query:
                    hay = " ".join([
                        str(row.get("room_name") or ""),
                        str(row.get("point_name") or ""),
                        str(row.get("slot_name") or ""),
                        str(row.get("notes") or ""),
                        str(row.get("recorded_by") or ""),
                    ]).lower()
                    if query not in hay:
                        continue

                session_key = self._session_key_from_row(row)
                bucket = grouped.get(session_key)
                if not bucket:
                    bucket = {
                        "session_key": session_key,
                        "timestamp": ts.strftime("%Y-%m-%d %H:%M:%S"),
                        "date": ts.strftime("%Y-%m-%d"),
                        "day": ts.strftime("%A"),
                        "time": ts.strftime("%H:%M"),
                        "room_name": str(row.get("room_name") or "").strip(),
                        "slot_name": str(row.get("slot_name") or "").strip(),
                        "notes": str(row.get("notes") or "").strip(),
                        "recorded_by": str(row.get("recorded_by") or "").strip(),
                        "point_temps": {},
                    }
                    grouped[session_key] = bucket
                bucket["point_temps"][str(row.get("point_name") or "").strip()] = self._safe_float(row.get("current_temp"))
                if not bucket.get("notes"):
                    bucket["notes"] = str(row.get("notes") or "").strip()
                if not bucket.get("recorded_by"):
                    bucket["recorded_by"] = str(row.get("recorded_by") or "").strip()
            except Exception:
                pass

        rows = []
        point_names = self.get_temperature_points()
        for bucket in grouped.values():
            temps = []
            row = {
                "session_key": bucket["session_key"],
                "timestamp": bucket["timestamp"],
                "date": bucket["date"],
                "day": bucket["day"],
                "time": bucket["time"],
                "room_name": bucket["room_name"],
                "slot_name": bucket["slot_name"],
                "notes": bucket["notes"],
                "recorded_by": bucket["recorded_by"],
            }
            for index, point_name in enumerate(point_names, start=1):
                value = bucket["point_temps"].get(point_name)
                row[f"p{index}"] = "" if value is None else format_temperature_display(value)
                if value is not None:
                    temps.append(value)
            if temps:
                row["min_temp"] = min(temps)
                row["max_temp"] = max(temps)
                row["current_temp"] = sum(temps) / len(temps)
            else:
                row["min_temp"] = None
                row["max_temp"] = None
                row["current_temp"] = None
            rows.append(row)

        rows.sort(key=lambda x: str(x.get("timestamp", "")), reverse=True)
        return rows

    def get_grouped_log_by_key(self, session_key):
        for row in self.get_grouped_logs():
            if row.get("session_key") == session_key:
                return row
        return None

    def _safe_float(self, value):
        try:
            return float(value)
        except Exception:
            return None

    def import_logs(self, file_path):
        ext = os.path.splitext(file_path or "")[1].lower()
        if ext in (".csv", ".txt"):
            return self._import_csv(file_path)
        if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            return self._import_xlsx(file_path)
        return False, "Unsupported file format", 0

    def _import_csv(self, file_path):
        added = 0
        try:
            with open(file_path, "r", newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    if self._import_row_dict(row):
                        added += 1
        except Exception as e:
            return False, str(e), 0
        if added:
            self.append_audit("temperature_log_imported", {"count": added, "file": os.path.basename(file_path)})
        return True, f"Imported {added} log(s)", added

    def _import_xlsx(self, file_path):
        try:
            import openpyxl
        except Exception:
            return False, "Excel import is not available in this build", 0
        added = 0
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            headers = [str(cell.value or "").strip() for cell in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                row = {}
                for i, h in enumerate(headers):
                    row[h] = r[i] if i < len(r) else None
                if self._import_row_dict(row):
                    added += 1
        except Exception as e:
            return False, str(e), 0
        if added:
            self.append_audit("temperature_log_imported", {"count": added, "file": os.path.basename(file_path)})
        return True, f"Imported {added} log(s)", added

    def _import_row_dict(self, row):
        if not isinstance(row, dict):
            return False
        keys = {str(k or "").strip().lower(): k for k in row.keys()}

        def pick(*names):
            for name in names:
                key = keys.get(name)
                if key is not None:
                    return row.get(key)
            return None

        ts = pick("timestamp", "date/time", "datetime", "date", "time")
        room = pick("room", "room name", "room_name", "location")
        slot = pick("slot", "slot name", "slot_name", "reading", "time slot")
        notes = pick("notes", "note", "reason")
        recorded_by = pick("user", "recorded by", "recorded_by")
        if room is None and ts is None:
            return False

        point_readings = {}
        point_names = self.get_temperature_points()
        for index, point_name in enumerate(point_names, start=1):
            value = pick(f"p{index}", point_name.lower(), point_name.lower().replace(" ", "_"))
            if value not in (None, ""):
                point_readings[point_name] = value

        if not point_readings:
            point = pick("point", "point name", "point_name", "probe", "sensor")
            cur = pick("current", "current temp", "current_temp", "temperature", "temp")
            if point and cur not in (None, ""):
                point_readings[str(point).strip()] = cur

        if not point_readings:
            return False

        ts_val = ts.strftime("%Y-%m-%d %H:%M:%S") if isinstance(ts, datetime) else ts
        room_val = room if room is not None else "Room"
        slot_val = str(slot or "").strip()
        if not slot_val:
            dt = self._parse_datetime(ts_val)
            slot_val = "Morning" if dt.hour < 14 else "Evening"
        try:
            ok, _, _ = self.add_session(ts_val, room_val, slot_val, point_readings, notes=notes or "", recorded_by=recorded_by or "")
            return ok
        except Exception:
            return False

    def export_csv(self, file_path, grouped_logs):
        rows = list(grouped_logs or [])
        try:
            with open(file_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["Date", "Day", "Time", "P1 (°C)", "P2 (°C)", "P3 (°C)", "P4 (°C)", "P5 (°C)", "P6 (°C)", "Slot", "Min (°C)", "Max (°C)", "Current (°C)", "User"])
                for row in rows:
                    writer.writerow(self._grouped_export_row(row))
            self.append_audit("temperature_log_exported_csv", {"file": os.path.basename(file_path), "count": len(rows)})
            return True, file_path
        except Exception as e:
            return False, str(e)

    def export_xlsx(self, file_path, grouped_logs):
        try:
            import openpyxl
        except Exception:
            return False, "Excel export is not available in this build"
        rows = list(grouped_logs or [])
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Temperature Log"
            ws.append(["Date", "Day", "Time", "P1 (°C)", "P2 (°C)", "P3 (°C)", "P4 (°C)", "P5 (°C)", "P6 (°C)", "Slot", "Min (°C)", "Max (°C)", "Current (°C)", "User"])
            for row in rows:
                ws.append(self._grouped_export_row(row))
            wb.save(file_path)
            self.append_audit("temperature_log_exported_xlsx", {"file": os.path.basename(file_path), "count": len(rows)})
            return True, file_path
        except Exception as e:
            return False, str(e)

    def _grouped_export_row(self, row):
        return [
            row.get("date", ""),
            row.get("day", ""),
            row.get("time", ""),
            row.get("p1", ""),
            row.get("p2", ""),
            row.get("p3", ""),
            row.get("p4", ""),
            row.get("p5", ""),
            row.get("p6", ""),
            row.get("slot_name", ""),
            self._fmt_temperature(row.get("min_temp")),
            self._fmt_temperature(row.get("max_temp")),
            self._fmt_temperature(row.get("current_temp")),
            row.get("recorded_by", ""),
        ]

    def _fmt_temperature(self, value):
        if value in (None, ""):
            return ""
        return format_temperature_display(value)

    def export_pdf(self, grouped_logs, title_suffix=""):
        try:
            from hope_pharma_complete import EnhancedPDFGenerator
        except Exception:
            return False, "PDF generator not available"
        out_folder = os.path.join(self.data_folder, "HopePharmaTemperatureLogs")
        try:
            os.makedirs(out_folder, exist_ok=True)
        except Exception:
            pass
        gen = EnhancedPDFGenerator(output_folder=out_folder)
        return gen.generate_temperature_log_pdf(grouped_logs, title_suffix=title_suffix)


class TemperatureLogScreen:
    def __init__(self, parent, invoice_manager, back_callback=None):
        self.parent = parent
        self.invoice_manager = invoice_manager
        self.back_callback = back_callback
        self.data_folder = getattr(invoice_manager, "invoice_folder", os.getcwd())
        self.manager = TemperatureLogManager(self.data_folder)
        self.filtered = []
        self._build_ui()
        self.refresh()

    def _build_ui(self):
        top = ttk.Frame(self.parent)
        top.pack(fill="x", pady=(0, 8))
        if self.back_callback:
            ttk.Button(top, text="← Back to Main Menu", command=self.back_callback).pack(side="left")
        ttk.Label(top, text="Temperature Log", font=("Helvetica", 20, "bold")).pack(side="left", padx=12)

        filters = ttk.LabelFrame(self.parent, text="Filter", padding=10)
        filters.pack(fill="x", pady=(0, 8))
        ttk.Label(filters, text="From (YYYY-MM-DD)").grid(row=0, column=0, sticky="w", padx=(0, 6))
        self.from_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.from_var, width=14).grid(row=0, column=1, sticky="w", padx=(0, 12))
        ttk.Label(filters, text="To (YYYY-MM-DD)").grid(row=0, column=2, sticky="w", padx=(0, 6))
        self.to_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.to_var, width=14).grid(row=0, column=3, sticky="w", padx=(0, 12))
        ttk.Label(filters, text="Search").grid(row=0, column=4, sticky="w", padx=(0, 6))
        self.room_var = tk.StringVar()
        ttk.Entry(filters, textvariable=self.room_var, width=24).grid(row=0, column=5, sticky="w", padx=(0, 12))
        ttk.Button(filters, text="Apply", command=self.refresh).grid(row=0, column=6, sticky="w")

        actions = ttk.Frame(self.parent)
        actions.pack(fill="x", pady=(0, 8))
        ttk.Button(actions, text="Add Multi-Point Reading", command=self.add_reading).pack(side="left")
        ttk.Button(actions, text="Edit Selected", command=self.edit_selected).pack(side="left", padx=6)
        ttk.Button(actions, text="Delete Selected", command=self.delete_selected).pack(side="left", padx=6)
        ttk.Button(actions, text="Import (Excel/CSV)", command=self.import_logs).pack(side="left", padx=6)
        ttk.Button(actions, text="Manage Points", command=self.manage_points).pack(side="left", padx=6)
        ttk.Button(actions, text="Export PDF", command=self.export_pdf).pack(side="right")
        ttk.Button(actions, text="Export Excel", command=self.export_excel).pack(side="right", padx=6)
        ttk.Button(actions, text="Export CSV", command=self.export_csv).pack(side="right", padx=6)

        table_frame = ttk.Frame(self.parent)
        table_frame.pack(fill="both", expand=True)
        cols = ("Date", "Day", "Time", "P1 (°C)", "P2 (°C)", "P3 (°C)", "P4 (°C)", "P5 (°C)", "P6 (°C)", "Slot", "Min (°C)", "Max (°C)", "Current (°C)", "User")
        self.tree = ttk.Treeview(table_frame, columns=cols, show="headings", height=18)
        for c in cols:
            self.tree.heading(c, text=c)
        self.tree.column("Date", width=100)
        self.tree.column("Day", width=100)
        self.tree.column("Time", width=80)
        for point_col in ("P1 (°C)", "P2 (°C)", "P3 (°C)", "P4 (°C)", "P5 (°C)", "P6 (°C)"):
            self.tree.column(point_col, width=70, anchor="e")
        self.tree.column("Slot", width=90)
        self.tree.column("Min (°C)", width=70, anchor="e")
        self.tree.column("Max (°C)", width=70, anchor="e")
        self.tree.column("Current (°C)", width=80, anchor="e")
        self.tree.column("User", width=120)
        y = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        x = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
        self.tree.pack(side="left", fill="both", expand=True)
        y.pack(side="right", fill="y")
        x.pack(side="bottom", fill="x")

    def refresh(self):
        self.filtered = self.manager.get_grouped_logs(
            self.from_var.get().strip(),
            self.to_var.get().strip(),
            self.room_var.get().strip(),
        )
        for item_id in self.tree.get_children():
            self.tree.delete(item_id)
        for row in self.filtered:
            self.tree.insert(
                "",
                "end",
                iid=row.get("session_key"),
                values=(
                    row.get("date", ""),
                    row.get("day", ""),
                    row.get("time", ""),
                    row.get("p1", ""),
                    row.get("p2", ""),
                    row.get("p3", ""),
                    row.get("p4", ""),
                    row.get("p5", ""),
                    row.get("p6", ""),
                    row.get("slot_name", ""),
                    self.manager._fmt_temperature(row.get("min_temp")),
                    self.manager._fmt_temperature(row.get("max_temp")),
                    self.manager._fmt_temperature(row.get("current_temp")),
                    row.get("recorded_by", ""),
                ),
            )

    def _get_selected_session_key(self):
        sel = self.tree.selection()
        return sel[0] if sel else None

    def add_reading(self):
        dlg = TemperatureLogBatchDialog(
            self.parent,
            title="Add Temperature Reading",
            point_values=self.manager.get_temperature_points(),
            slot_values=self.manager.get_temperature_slots(),
        )
        self.parent.wait_window(dlg.dialog)
        if not dlg.result:
            return
        date_key = self.manager.get_date_key(dlg.result.get("timestamp"))
        room = str(dlg.result.get("room_name") or "").strip()
        slot = str(dlg.result.get("slot_name") or "").strip()
        if self.manager.session_exists(date_key, room, slot):
            if not messagebox.askyesno("Warning", "A session for this room/date/slot already exists.\nDo you want to save anyway?"):
                return
        try:
            ok, msg, _ = self.manager.add_session(**dlg.result)
        except Exception as exc:
            messagebox.showerror("Error", str(exc))
            return
        if not ok:
            messagebox.showerror("Error", msg)
            return
        self.refresh()

    def edit_selected(self):
        session_key = self._get_selected_session_key()
        if not session_key:
            messagebox.showwarning("Warning", "Select a row first")
            return
        current = self.manager.get_grouped_log_by_key(session_key)
        if not current:
            messagebox.showerror("Error", "Temperature log session not found")
            return
        dlg = TemperatureLogBatchDialog(
            self.parent,
            title="Edit Temperature Reading",
            initial=current,
            point_values=self.manager.get_temperature_points(),
            slot_values=self.manager.get_temperature_slots(),
        )
        self.parent.wait_window(dlg.dialog)
        if not dlg.result:
            return
        date_key = self.manager.get_date_key(dlg.result.get("timestamp"))
        room = str(dlg.result.get("room_name") or "").strip()
        slot = str(dlg.result.get("slot_name") or "").strip()
        if self.manager.session_exists(date_key, room, slot, exclude_session_key=session_key):
            if not messagebox.askyesno("Warning", "A session for this room/date/slot already exists.\nDo you want to save anyway?"):
                return
        try:
            ok, msg, _ = self.manager.update_session(session_key=session_key, **dlg.result)
        except Exception as exc:
            messagebox.showerror("Error", str(exc))
            return
        if not ok:
            messagebox.showerror("Error", msg)
            return
        self.refresh()

    def delete_selected(self):
        session_key = self._get_selected_session_key()
        if not session_key:
            return
        if not messagebox.askyesno("Confirm", "Delete the selected temperature log session?"):
            return
        ok, msg = self.manager.delete_session(session_key)
        if not ok:
            messagebox.showerror("Error", msg)
            return
        self.refresh()

    def import_logs(self):
        path = filedialog.askopenfilename(
            title="Import Temperature Log",
            filetypes=[
                ("Excel Files", "*.xlsx;*.xlsm;*.xltx;*.xltm"),
                ("CSV Files", "*.csv"),
                ("All Files", "*.*"),
            ],
        )
        if not path:
            return
        ok, msg, _ = self.manager.import_logs(path)
        if not ok:
            messagebox.showerror("Error", msg)
            return
        messagebox.showinfo("Success", msg)
        self.refresh()

    def export_csv(self):
        path = filedialog.asksaveasfilename(
            title="Export Temperature Log (CSV)",
            defaultextension=".csv",
            filetypes=[("CSV Files", "*.csv")],
        )
        if not path:
            return
        ok, msg = self.manager.export_csv(path, self.filtered)
        if not ok:
            messagebox.showerror("Error", msg)
            return
        messagebox.showinfo("Success", f"Saved:\n{path}")

    def export_excel(self):
        path = filedialog.asksaveasfilename(
            title="Export Temperature Log (Excel)",
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
        )
        if not path:
            return
        ok, msg = self.manager.export_xlsx(path, self.filtered)
        if not ok:
            messagebox.showerror("Error", msg)
            return
        messagebox.showinfo("Success", f"Saved:\n{path}")

    def export_pdf(self):
        title_suffix = ""
        if self.from_var.get().strip() or self.to_var.get().strip() or self.room_var.get().strip():
            parts = []
            if self.from_var.get().strip():
                parts.append(f"From {self.from_var.get().strip()}")
            if self.to_var.get().strip():
                parts.append(f"To {self.to_var.get().strip()}")
            if self.room_var.get().strip():
                parts.append(self.room_var.get().strip())
            title_suffix = " - " + " | ".join(parts)
        ok, res = self.manager.export_pdf(self.filtered, title_suffix=title_suffix)
        if not ok:
            messagebox.showerror("Error", res)
            return
        try:
            import webbrowser
            webbrowser.open("file://" + os.path.abspath(res))
        except Exception:
            pass
        messagebox.showinfo("Success", f"PDF saved:\n{res}")

    def manage_points(self):
        points = self.manager.get_temperature_points()
        dlg = FixedListEditorDialog(
            self.parent,
            title="Edit Temperature Points",
            labels=[f"P{i}" for i in range(1, 7)],
            values=points[:6],
        )
        self.parent.wait_window(dlg.dialog)
        if dlg.result is None:
            return
        ok, msg = self.manager.set_temperature_points(dlg.result)
        if not ok:
            messagebox.showerror("Error", msg)
            return
        self.refresh()


class TemperatureLogBatchDialog:
    def __init__(self, parent, title="Temperature Reading", initial=None, point_values=None, slot_values=None):
        self.result = None
        self.point_values = list(point_values or [])
        self.slot_values = list(slot_values or [])
        self.dialog = tk.Toplevel(parent)
        self.dialog.title(title)
        self.dialog.geometry("760x650")
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self._build(initial or {})

    def _build(self, initial):
        frm = ttk.Frame(self.dialog, padding=12)
        frm.pack(fill="both", expand=True)

        ttk.Label(frm, text="Date/Time (YYYY-MM-DD HH:MM)").grid(row=0, column=0, sticky="w", pady=6)
        self.ts_var = tk.StringVar(value=str(initial.get("timestamp") or datetime.now().strftime("%Y-%m-%d %H:%M")))
        ttk.Entry(frm, textvariable=self.ts_var, width=24).grid(row=0, column=1, sticky="w", pady=6, padx=(0, 16))

        ttk.Label(frm, text="Room Name *").grid(row=0, column=2, sticky="w", pady=6)
        self.room_var = tk.StringVar(value=str(initial.get("room_name") or ""))
        ttk.Entry(frm, textvariable=self.room_var, width=24).grid(row=0, column=3, sticky="w", pady=6)

        ttk.Label(frm, text="Slot *").grid(row=1, column=0, sticky="w", pady=6)
        default_slot = str(initial.get("slot_name") or "")
        if not default_slot and self.slot_values:
            default_slot = self.slot_values[0]
        self.slot_var = tk.StringVar(value=default_slot)
        ttk.Combobox(frm, textvariable=self.slot_var, values=self.slot_values, width=22, state="readonly").grid(row=1, column=1, sticky="w", pady=6)

        ttk.Label(frm, text="User *").grid(row=1, column=2, sticky="w", pady=6)
        default_user = str(initial.get("recorded_by") or "")
        if not default_user:
            try:
                default_user = getpass.getuser()
            except Exception:
                default_user = ""
        self.user_var = tk.StringVar(value=default_user)
        ttk.Entry(frm, textvariable=self.user_var, width=24).grid(row=1, column=3, sticky="w", pady=6)

        summary = ttk.LabelFrame(frm, text="Summary", padding=10)
        summary.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(6, 10))
        summary.columnconfigure(1, weight=1)
        self.min_label = tk.StringVar(value="Min: - °C")
        self.max_label = tk.StringVar(value="Max: - °C")
        self.current_label = tk.StringVar(value="Current: - °C")
        ttk.Label(summary, textvariable=self.min_label).grid(row=0, column=0, sticky="w", padx=(0, 16))
        ttk.Label(summary, textvariable=self.max_label).grid(row=0, column=1, sticky="w", padx=(0, 16))
        ttk.Label(summary, textvariable=self.current_label).grid(row=0, column=2, sticky="w")

        points_frame = ttk.LabelFrame(frm, text="Point Readings", padding=10)
        points_frame.grid(row=3, column=0, columnspan=4, sticky="nsew")
        frm.rowconfigure(3, weight=1)
        for col in range(4):
            points_frame.columnconfigure(col, weight=1)

        self.point_vars = {}
        self.point_entries = {}
        for index, point_name in enumerate(self.point_values, start=1):
            row_idx = (index - 1) // 2
            col_idx = ((index - 1) % 2) * 2
            ttk.Label(points_frame, text=f"P{index} (°C) - {point_name} *").grid(row=row_idx, column=col_idx, sticky="w", pady=6, padx=(0, 8))
            value = format_temperature_display(initial.get(f"p{index}") or "")
            var = tk.StringVar(value=value)
            var.trace_add("write", self._update_summary)
            entry = ttk.Entry(points_frame, textvariable=var, width=16)
            entry.grid(row=row_idx, column=col_idx + 1, sticky="w", pady=6, padx=(0, 16))
            entry.bind("<FocusOut>", lambda _event, v=var: self._format_temperature_var(v))
            self.point_vars[point_name] = var
            self.point_entries[point_name] = entry

        ttk.Label(frm, text="Notes").grid(row=4, column=0, sticky="nw", pady=8)
        self.notes = scrolledtext.ScrolledText(frm, height=5, width=60)
        self.notes.grid(row=4, column=1, columnspan=3, sticky="ew", pady=8)
        self.notes.insert("1.0", str(initial.get("notes") or ""))

        btns = ttk.Frame(frm)
        btns.grid(row=5, column=0, columnspan=4, sticky="w", pady=(12, 0))
        ttk.Button(btns, text="Save", command=self._save).pack(side="left")
        ttk.Button(btns, text="Cancel", command=self.dialog.destroy).pack(side="left", padx=6)

        self._update_summary()

    def _format_temperature_var(self, var):
        raw = str(var.get() or "").strip()
        if not raw:
            return
        try:
            formatted = format_temperature_display(raw)
        except Exception:
            return
        if formatted and raw != formatted:
            var.set(formatted)

    def _update_summary(self, *_args):
        values = []
        for var in self.point_vars.values():
            raw = str(var.get() or "").strip()
            if not raw:
                continue
            try:
                values.append(parse_temperature_value(raw))
            except Exception:
                continue
        if not values:
            self.min_label.set("Min: - °C")
            self.max_label.set("Max: - °C")
            self.current_label.set("Current: - °C")
            return
        self.min_label.set(f"Min: {format_temperature_display(min(values))}")
        self.max_label.set(f"Max: {format_temperature_display(max(values))}")
        self.current_label.set(f"Current: {format_temperature_display(sum(values) / len(values))}")

    def _save(self):
        ts = self.ts_var.get().strip()
        room = self.room_var.get().strip()
        slot = self.slot_var.get().strip()
        recorded_by = self.user_var.get().strip()
        notes = self.notes.get("1.0", tk.END).strip()
        if not room:
            messagebox.showwarning("Warning", "Room name is required")
            return
        if not slot:
            messagebox.showwarning("Warning", "Slot is required")
            return
        if not recorded_by:
            messagebox.showwarning("Warning", "User is required")
            return

        point_readings = {}
        try:
            for point_name, var in self.point_vars.items():
                raw = str(var.get() or "").strip()
                if not raw:
                    messagebox.showwarning("Warning", f"{point_name} temperature is required")
                    return
                point_readings[point_name] = parse_temperature_value(raw, f"{point_name} temperature")
                var.set(format_temperature_display(point_readings[point_name]))
        except Exception as exc:
            messagebox.showwarning("Warning", str(exc) or "All point temperatures must be valid")
            return

        self.result = {
            "timestamp": ts,
            "room_name": room,
            "slot_name": slot,
            "point_readings": point_readings,
            "recorded_by": recorded_by,
            "notes": notes,
        }
        self.dialog.destroy()


class FixedListEditorDialog:
    def __init__(self, parent, title, labels, values):
        self.result = None
        self.dialog = tk.Toplevel(parent)
        self.dialog.title(title)
        self.dialog.geometry("420x320")
        self.dialog.transient(parent)
        self.dialog.grab_set()

        frm = ttk.Frame(self.dialog, padding=12)
        frm.pack(fill="both", expand=True)

        self.vars = []
        labels = list(labels or [])
        values = list(values or [])
        for i, lbl in enumerate(labels):
            ttk.Label(frm, text=str(lbl)).grid(row=i, column=0, sticky="w", pady=6, padx=(0, 8))
            var = tk.StringVar(value=str(values[i] if i < len(values) else "").strip())
            ttk.Entry(frm, textvariable=var, width=28).grid(row=i, column=1, sticky="w", pady=6)
            self.vars.append(var)

        btns = ttk.Frame(frm)
        btns.grid(row=len(labels) + 1, column=0, columnspan=2, sticky="w", pady=(12, 0))
        ttk.Button(btns, text="Save", command=self._save).pack(side="left")
        ttk.Button(btns, text="Cancel", command=self.dialog.destroy).pack(side="left", padx=6)

    def _save(self):
        out = [str(v.get() or "").strip() for v in (self.vars or [])]
        if any(not x for x in out):
            messagebox.showwarning("Warning", "All fields are required")
            return
        self.result = out
        self.dialog.destroy()
